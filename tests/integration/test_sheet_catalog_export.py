"""图纸目录导出执行动作与 Artifact 登记集成测试（PLAN-DM-020 Task 9 / SPEC-DM-012 §8.2、§9、§10）。

覆盖：预览后执行成功并回读 XLSX/Artifact、未保存模板可执行、修订/模板/扩展
版本/动作摘要变化 ``REPREVIEW_REQUIRED``、无效/漂移/复用授权拒绝、候选失败不
登记、成功响应无后台字段、Artifact 三态可用性、桥与 API runtime 共享同一
``SaveGrantStore``，以及整条预览→执行链路的只读不变量；PLAN-DM-025 Task 4 的
设置快照绑定：过滤投影同时作用于预览与候选行、全部过滤导出只有表头、预览后
设置变化在候选目录与授权消费之前以 ``EXTENSION_SETTINGS_CHANGED``/409 拒绝，
以及预览末尾 best-effort 偏好写入不使该预览自行过期。
"""

import hashlib
import uuid
from copy import deepcopy
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

import dst_manager.extensions.builtin.sheet_catalog.extension as extension_module
import dst_manager.extensions.registry as registry_module
from dst_manager.application.extensions.runtime import ExtensionRuntime
from dst_manager.config import Settings
from dst_manager.extensions.builtin.sheet_catalog.templates import DEFAULT_TEMPLATE
from dst_manager.extensions.registry import ExtensionRegistry
from dst_manager.extensions.save_grants import SaveGrantStore
from dst_manager.infrastructure.acsm_xml import AcsmDocument
from dst_manager.infrastructure.dst_codec import DstCodec
from dst_manager.infrastructure.extension_workspace import ExtensionWorkspaceReader
from dst_manager.infrastructure.persistence.database import Database
from dst_manager.infrastructure.persistence.extensions import ExtensionStore
from dst_manager.interfaces.api import create_app
from dst_manager.interfaces.shell import ShellBridge

SHEET_CATALOG_ID = "dst-manager.sheet-catalog"
ACTION_ID = "export-xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_client(tmp_path, save_grants=None, **kwargs) -> TestClient:
    return TestClient(
        create_app(
            Settings(data_dir=tmp_path / "data"), save_grants=save_grants, **kwargs
        )
    )


def grant_store(client: TestClient) -> SaveGrantStore:
    """API runtime 持有的授权存储（桌面装配注入同一实例）。"""
    return client.app.state.extension_runtime.save_grants


def candidate_files(client: TestClient) -> list[Path]:
    """应用临时候选目录的残留文件（成功/失败后都必须为空）。"""
    root = Path(client.app.state.extension_runtime.proposal_root)
    return sorted(p for p in root.rglob("*") if p.is_file()) if root.is_dir() else []


def template_payload(template=DEFAULT_TEMPLATE, template_id=None) -> dict:
    return {
        "template_id": str(template_id) if template_id else None,
        "name": template.name,
        "schema_version": template.schema_version,
        "columns": [
            {
                "column_id": str(column.column_id),
                "header": column.header,
                "expression": column.expression,
            }
            for column in template.columns
        ],
    }


def draft_template() -> dict:
    """未保存草稿模板：自定义表头与表达式（template_id=null）。"""
    return {
        "template_id": None,
        "name": "未保存草稿",
        "schema_version": 1,
        "columns": [
            {
                "column_id": str(uuid.uuid5(uuid.NAMESPACE_URL, "export:number")),
                "header": "编号",
                "expression": "{sheet.number}",
            }
        ],
    }


def format_code_template(expression: str) -> dict:
    """带数字格式码的未保存草稿模板（SPEC-DM-012 §5.4）。"""
    return {
        "template_id": None,
        "name": "格式码草稿",
        "schema_version": 1,
        "columns": [
            {
                "column_id": str(uuid.uuid5(uuid.NAMESPACE_URL, "export:format-code")),
                "header": "图号",
                "expression": expression,
            }
        ],
    }


def add_excluded_sheet(dst: Path, *, number: str, title: str) -> None:
    """在测试工作区追加一张图纸（复制首张节点并改写 ID/编号/图名与布局名）。"""
    document = AcsmDocument(DstCodec().decode_file(dst))
    subset = document.root.xpath("//*[local-name()='AcSmSubset']")[0]
    first = subset.xpath("./*[local-name()='AcSmSheet']")[0]
    second = deepcopy(first)
    for index, node in enumerate([second, *second.xpath(".//*[@ID]")], start=20):
        node.set("ID", f"g00000000-0000-0000-0000-{index:012X}")
    second.xpath("./*[local-name()='AcSmProp' and @propname='Number']")[0].text = number
    second.xpath("./*[local-name()='AcSmProp' and @propname='Title']")[0].text = title
    second.xpath(
        "./*[local-name()='AcSmAcDbLayoutReference']"
        "/*[local-name()='AcSmProp' and @propname='Name']",
    )[0].text = f"{number} {title}"
    subset.append(second)
    DstCodec().encode_file(document.to_bytes(), dst)


def open_workspace(client: TestClient, tiny_workspace) -> dict:
    opened = client.post("/api/workspaces/open", json={"dst_path": str(tiny_workspace[0])})
    assert opened.status_code == 200
    return opened.json()


def preview_action(client: TestClient, workspace: dict, template=None) -> dict:
    body = template if template is not None else template_payload()
    resp = client.post(
        f"/api/extensions/{SHEET_CATALOG_ID}/actions/{ACTION_ID}/preview",
        json={
            "workspace_id": workspace["id"],
            "base_revision_id": workspace["revision_id"],
            "template": body,
        },
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def create_grant(store: SaveGrantStore, workspace: dict, target: Path):
    return store.create(SHEET_CATALOG_ID, ACTION_ID, workspace["id"], target)


def execute_action(
    client: TestClient,
    workspace: dict,
    preview_body: dict,
    grant_id: str,
    template=None,
    settings_revision="unset",
):
    body = template if template is not None else template_payload()
    return client.post(
        f"/api/extensions/{SHEET_CATALOG_ID}/actions/{ACTION_ID}/execute",
        json={
            "workspace_id": workspace["id"],
            "base_revision_id": workspace["revision_id"],
            "template": body,
            "preview_digest": preview_body["preview_digest"],
            "save_grant_id": grant_id,
            # 默认原样重复提交预览绑定的设置修订（前端契约）；用例可显式提交过期值
            "settings_revision": (
                preview_body.get("settings_revision", 0)
                if settings_revision == "unset"
                else settings_revision
            ),
        },
    )


def artifact_count(client: TestClient) -> int:
    """Artifact 登记数（漂移拒绝必须不留下成功登记）。"""
    with client.app.state.service.database.sessions() as session:
        return session.execute(text("SELECT COUNT(*) FROM artifacts")).scalar_one()


def assert_error_contract(body: dict) -> None:
    assert set(body) == {"code", "message_key", "params", "message"}
    assert body["code"] and body["message_key"] and body["message"]
    assert isinstance(body["params"], dict)


# ---------------------------------------------------------------------------
# 预览后执行成功：XLSX 回读、Artifact 登记与查询
# ---------------------------------------------------------------------------


def test_preview_then_execute_roundtrip_registers_artifact(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    target = tmp_path / "exports"
    target.mkdir()
    grant = create_grant(grant_store(client), workspace, target / "图纸目录.xlsx")
    preview_body = preview_action(client, workspace)

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 200, resp.text
    body = resp.json()
    # 成功响应只含 SPEC §8.2 四字段：不含 sha256/source_revision/extension_version
    assert set(body) == {"artifact_id", "file_name", "output_path", "warnings"}
    assert body["file_name"] == "图纸目录.xlsx"
    assert Path(body["output_path"]).is_file()
    assert body["warnings"] == []

    # openpyxl 回读真实落盘文件（SPEC §9）
    workbook = load_workbook(body["output_path"])
    assert workbook.sheetnames == ["图纸目录"]
    worksheet = workbook["图纸目录"]
    assert [cell.value for cell in worksheet[1]] == ["图号", "图名", "文件名"]
    assert [[cell.value for cell in row] for row in worksheet.iter_rows(min_row=2)] == [
        ["001", "平面", "A.dwg"]
    ]
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref is not None
    workbook.close()

    # Artifact 查询含后台字段与可用性派生
    artifact = client.get(f"/api/artifacts/{body['artifact_id']}")
    assert artifact.status_code == 200
    record = artifact.json()
    assert record["extension_id"] == SHEET_CATALOG_ID
    assert record["extension_version"] == "0.1.0"
    assert record["workspace_id"] == workspace["id"]
    assert record["source_revision_id"] == workspace["revision_id"]
    assert record["kind"] == "sheet-catalog"
    assert record["media_type"] == XLSX_MEDIA_TYPE
    assert record["management_relation"] == "external"
    assert record["file_name"] == "图纸目录.xlsx"
    assert record["availability"] == "AVAILABLE"
    saved = Path(body["output_path"])
    assert record["size_bytes"] == saved.stat().st_size
    assert record["sha256"] == hashlib.sha256(saved.read_bytes()).hexdigest()

    # 应用临时候选目录清空（SPEC §10）
    assert candidate_files(client) == []


def test_unsaved_draft_template_executes_without_server_state(tmp_path, tiny_workspace):
    """模板来自请求快照：未保存草稿（template_id=null）也能执行。"""
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "草稿.xlsx")
    template = draft_template()
    preview_body = preview_action(client, workspace, template)

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id, template)

    assert resp.status_code == 200, resp.text
    body = resp.json()
    workbook = load_workbook(body["output_path"])
    worksheet = workbook["图纸目录"]
    assert [cell.value for cell in worksheet[1]] == ["编号"]
    assert worksheet["A2"].value == "001"
    workbook.close()
    assert candidate_files(client) == []


def test_number_format_code_exports_string_cell_with_leading_zeros(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    template = format_code_template("{sheet.number:0000}")
    target = tmp_path / "exports"
    target.mkdir()
    grant = create_grant(grant_store(client), workspace, target / "带格式码.xlsx")
    preview_body = preview_action(client, workspace, template=template)

    assert preview_body["executable"] is True
    assert preview_body["rows"] == [["0001"]]

    resp = execute_action(
        client, workspace, preview_body, grant.save_grant_id, template=template
    )
    assert resp.status_code == 200, resp.text

    workbook = load_workbook(resp.json()["output_path"])
    worksheet = workbook[workbook.sheetnames[0]]
    cell = worksheet.cell(row=2, column=1)
    assert cell.value == "0001"
    assert cell.data_type == "s"  # 文本单元格，不是数值
    workbook.close()


# ---------------------------------------------------------------------------
# REPREVIEW_REQUIRED：修订 / 模板 / 扩展版本 / 动作摘要变化
# ---------------------------------------------------------------------------


def test_revision_change_requires_repreview_and_keeps_grant(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    stale = dict(workspace)
    stale["revision_id"] = "f" * 64

    resp = execute_action(client, stale, preview_body, grant.save_grant_id)

    assert resp.status_code == 409
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "REPREVIEW_REQUIRED"
    assert body["message_key"] == "errors.extension.repreviewRequired"
    # 授权未被烧毁：刷新预览后可用同一授权重试（保留编辑语义）
    assert grant_store(client).active_count() == 1
    retry = execute_action(client, workspace, preview_body, grant.save_grant_id)
    assert retry.status_code == 200


def test_template_change_requires_repreview(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    changed = draft_template()  # 与预览所用模板不同列

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id, changed)

    assert resp.status_code == 409
    assert resp.json()["code"] == "REPREVIEW_REQUIRED"


def test_number_format_code_change_requires_repreview(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "带格式码.xlsx")
    # 预览用无格式码表达式，导出改用只差格式码的模板：
    # 列 ID 与表头完全相同，只有 token 的格式宽度不同。
    preview_body = preview_action(client, workspace, template=format_code_template("{sheet.number}"))
    changed = format_code_template("{sheet.number:0000}")

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id, template=changed)

    assert resp.status_code == 409
    assert resp.json()["code"] == "REPREVIEW_REQUIRED"


def test_extension_version_change_requires_repreview(tmp_path, tiny_workspace, monkeypatch):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    monkeypatch.setattr(extension_module, "EXTENSION_VERSION", "9.9.9")

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 409
    assert resp.json()["code"] == "REPREVIEW_REQUIRED"


def test_action_summary_change_requires_repreview(tmp_path, tiny_workspace, monkeypatch):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    monkeypatch.setattr(extension_module, "PREVIEW_ACTION_ID", "export-xlsx-v2")

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 409
    assert resp.json()["code"] == "REPREVIEW_REQUIRED"


def test_capability_unavailable_returns_503_and_keeps_grant(tmp_path, tiny_workspace, monkeypatch):
    # 能力不可发放（WAITING_DEPENDENCY）：执行不得编造结果；预览先于执行不可用
    monkeypatch.setattr(registry_module, "AVAILABLE_CAPABILITIES", frozenset())
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")

    resp = execute_action(client, workspace, {"preview_digest": "0" * 64}, grant.save_grant_id)

    assert resp.status_code == 503
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "EXTENSION_CAPABILITY_UNAVAILABLE"
    assert grant_store(client).active_count() == 1


def test_disabled_extension_rejects_execute(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    assert client.patch(
        f"/api/extensions/{SHEET_CATALOG_ID}/state", json={"enabled": False}
    ).status_code == 200

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 409
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "EXTENSION_DISABLED"


# ---------------------------------------------------------------------------
# 设置快照绑定：过滤投影与漂移拒绝（PLAN-DM-025 Task 4 / ARCH-DM-006 §11）
# ---------------------------------------------------------------------------


def test_filter_settings_project_preview_and_candidate_rows(tmp_path, tiny_workspace):
    """预览与候选行从同一设置快照读取过滤词，计数与导出内容一致（部分过滤）。"""
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    add_excluded_sheet(tiny_workspace[0], number="002", title="作废-平面")
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "过滤.xlsx")

    saved = client.put(
        f"/api/extensions/{SHEET_CATALOG_ID}/settings",
        json={
            "schema_version": 2,
            "expected_revision": 0,
            "value": {"excluded_title_keywords": " 作废，DRAFT "},
        },
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["value"]["excluded_title_keywords"] == ["作废", "DRAFT"]

    preview_body = preview_action(client, workspace)

    assert preview_body["settings_revision"] == saved.json()["revision"]
    assert preview_body["executable"] is True
    assert preview_body["rows"] == [["001", "平面", "A.dwg"]]
    assert (preview_body["total_rows"], preview_body["filtered_rows"]) == (1, 1)

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 200, resp.text
    workbook = load_workbook(resp.json()["output_path"])
    worksheet = workbook["图纸目录"]
    assert [cell.value for cell in worksheet[1]] == ["图号", "图名", "文件名"]
    assert [[cell.value for cell in row] for row in worksheet.iter_rows(min_row=2)] == [
        ["001", "平面", "A.dwg"]
    ]
    workbook.close()
    assert candidate_files(client) == []


def test_filter_everything_away_exports_header_only_xlsx(tmp_path, tiny_workspace):
    """全部过滤仍可执行：登记 Artifact 且候选工作簿只有表头。"""
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "空目录.xlsx")
    saved = client.put(
        f"/api/extensions/{SHEET_CATALOG_ID}/settings",
        json={
            "schema_version": 2,
            "expected_revision": 0,
            "value": {"excluded_title_keywords": ["平面"]},
        },
    )
    assert saved.status_code == 200, saved.text

    preview_body = preview_action(client, workspace)

    assert preview_body["executable"] is True
    assert preview_body["errors"] == []
    assert preview_body["rows"] == []
    assert (preview_body["total_rows"], preview_body["filtered_rows"]) == (0, 1)

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 200, resp.text
    workbook = load_workbook(resp.json()["output_path"])
    worksheet = workbook["图纸目录"]
    assert [cell.value for cell in worksheet[1]] == ["图号", "图名", "文件名"]
    assert worksheet.max_row == 1
    assert list(worksheet.iter_rows(min_row=2)) == []
    workbook.close()
    artifact = client.get(f"/api/artifacts/{resp.json()['artifact_id']}")
    assert artifact.status_code == 200
    assert artifact.json()["availability"] == "AVAILABLE"


def test_settings_change_after_preview_rejects_execute_before_side_effects(
    tmp_path, tiny_workspace
):
    """预览后设置变化：候选目录与授权消费之前返回 EXTENSION_SETTINGS_CHANGED/409。"""
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    target = exports / "图纸目录.xlsx"
    grant = create_grant(grant_store(client), workspace, target)
    preview_body = preview_action(client, workspace)
    saved = client.put(
        f"/api/extensions/{SHEET_CATALOG_ID}/settings",
        json={
            "schema_version": 2,
            "expected_revision": preview_body["settings_revision"],
            "value": {"excluded_title_keywords": ["作废"]},
        },
    )
    assert saved.status_code == 200, saved.text

    response = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert response.status_code == 409
    body = response.json()
    assert_error_contract(body)
    assert body["code"] == "EXTENSION_SETTINGS_CHANGED"
    assert body["message_key"] == "errors.extension.settingsChanged"
    assert body["params"]["expected_revision"] == preview_body["settings_revision"]
    assert body["params"]["current_revision"] == saved.json()["revision"]
    # 漂移必须在候选目录分配与授权消费之前拒绝：无候选文件、无目标文件、无 Artifact
    assert candidate_files(client) == []
    assert not target.exists()
    assert artifact_count(client) == 0
    # 授权未被烧毁：重新预览（携带新设置修订）后同一授权仍可使用
    assert grant_store(client).active_count() == 1
    refreshed = preview_action(client, workspace)
    assert refreshed["settings_revision"] == saved.json()["revision"]
    retried = execute_action(client, workspace, refreshed, grant.save_grant_id)
    assert retried.status_code == 200, retried.text


def test_stale_settings_revision_from_client_rejects_execute(tmp_path, tiny_workspace):
    """客户端重复提交过期设置修订（预览后外部修改设置的另一路径）同样 409。"""
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)

    response = execute_action(
        client, workspace, preview_body, grant.save_grant_id, settings_revision=99
    )

    assert response.status_code == 409
    assert response.json()["code"] == "EXTENSION_SETTINGS_CHANGED"
    assert artifact_count(client) == 0


def test_higher_schema_settings_fail_closed_for_preview_and_execute(tmp_path, tiny_workspace):
    """已存设置 Schema 高于当前版本：动作调用 fail-closed，不按当前语义执行。

    ARCH-DM-006 §8.1/§11：未知高版本只能只读保留；Runtime 在创建上下文前取得
    快照失败时必须拒绝动作（预览 409 与执行 409 同一稳定码），而不是当成空配置
    或 Provider 默认值。
    """
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    ExtensionStore(client.app.state.service.database.sessions).put_settings(
        SHEET_CATALOG_ID, 3, {"future_field": 1}, 0
    )
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    target = exports / "out.xlsx"
    grant = create_grant(grant_store(client), workspace, target)

    preview = client.post(
        f"/api/extensions/{SHEET_CATALOG_ID}/actions/{ACTION_ID}/preview",
        json={
            "workspace_id": workspace["id"],
            "base_revision_id": workspace["revision_id"],
            "template": template_payload(),
        },
    )
    assert preview.status_code == 409
    assert preview.json()["code"] == "EXTENSION_SETTINGS_SCHEMA_NEWER"

    response = execute_action(
        client, workspace, {"preview_digest": "0" * 64}, grant.save_grant_id
    )

    assert response.status_code == 409
    assert response.json()["code"] == "EXTENSION_SETTINGS_SCHEMA_NEWER"
    assert candidate_files(client) == []
    assert not target.exists()
    assert artifact_count(client) == 0
    assert grant_store(client).active_count() == 1


def test_workspace_preference_write_never_expires_the_same_preview(tmp_path, tiny_workspace):
    """预览末尾的 best-effort“上次选中模板”偏好写入不得使该预览自行过期。

    ARCH-DM-006 §11（MEMO-DM-033 已修订）：只有会改变动作输出且未进入规范化
    动作请求的工作区偏好才绑定其修订。当前唯一偏好是“上次选中的已保存模板
    ID”，模板身份与内容已进入规范化动作请求，因此它不绑定；若绑定，每次预览
    都会被自己的偏好写入立即作废。未来新增任何会改变输出且未进入规范化动作
    请求的偏好时，必须把其修订绑定进 ``preview_digest``，并以
    ``REPREVIEW_REQUIRED``/409 拒绝漂移。
    """
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    template_id = uuid.uuid5(uuid.NAMESPACE_URL, "export:preference-template")
    template = template_payload(template_id=template_id)

    preview_body = preview_action(client, workspace, template)

    preference = client.get(
        f"/api/extensions/{SHEET_CATALOG_ID}/workspaces/{workspace['id']}/preferences"
    ).json()
    assert preference["value"] == {"template_id": str(template_id)}
    # 偏好词汇目前只有唯一一项非输出绑定项；新增输出相关偏好必须重定绑定规则
    assert set(preference["value"]) == {"template_id"}

    resp = execute_action(
        client, workspace, preview_body, grant.save_grant_id, template=template
    )

    assert resp.status_code == 200, resp.text


# ---------------------------------------------------------------------------
# 保存授权：无效 / 漂移 / 复用
# ---------------------------------------------------------------------------


def test_unknown_grant_rejected_with_save_grant_invalid(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    preview_body = preview_action(client, workspace)

    resp = execute_action(client, workspace, preview_body, "no-such-grant")

    assert resp.status_code == 409
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "SAVE_GRANT_INVALID"
    assert body["message_key"] == "errors.extension.saveGrantInvalid"


def test_mismatched_workspace_grant_rejected(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = SaveGrantStore().create(SHEET_CATALOG_ID, ACTION_ID, "ws-other", exports / "out.xlsx")
    preview_body = preview_action(client, workspace)

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 409
    assert resp.json()["code"] == "SAVE_GRANT_INVALID"


def test_drifted_destination_rejected_with_export_destination_changed(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    (exports / "out.xlsx").write_bytes(b"created after grant")

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 409
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "EXPORT_DESTINATION_CHANGED"
    assert body["message_key"] == "errors.extension.exportDestinationChanged"
    assert not candidate_files(client)


def test_grant_cannot_be_reused_after_successful_execute(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    assert execute_action(client, workspace, preview_body, grant.save_grant_id).status_code == 200

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 409
    assert resp.json()["code"] == "SAVE_GRANT_INVALID"


# ---------------------------------------------------------------------------
# 候选失败：不登记 Artifact、不留半文件
# ---------------------------------------------------------------------------


def test_invalid_candidate_fails_without_artifact(tmp_path, tiny_workspace, monkeypatch):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)

    def broken_write_candidate(path, headers, rows):
        path.write_bytes(b"not an xlsx")
        from dst_manager.extensions.builtin.sheet_catalog.workbook import (
            WorkbookSummary,
        )

        return WorkbookSummary(worksheet_name="图纸目录", headers=tuple(headers), data_rows=len(rows))

    monkeypatch.setattr(extension_module, "write_candidate", broken_write_candidate)

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 503
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "ARTIFACT_WRITE_FAILED"
    assert body["message_key"] == "errors.extension.artifactWriteFailed"
    # 目标未被写入、无半文件残留、无 Artifact 登记
    assert not (exports / "out.xlsx").exists()
    assert list(exports.iterdir()) == []
    assert candidate_files(client) == []
    with client.app.state.service.database.sessions() as session:
        assert session.execute(text("SELECT COUNT(*) FROM artifacts")).scalar_one() == 0


def test_candidate_directory_os_error_is_contract_artifact_write_failed(
    tmp_path, tiny_workspace, monkeypatch
):
    """候选目录创建的 OS 级失败契约化为 ARTIFACT_WRITE_FAILED，授权未被消费。"""
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    proposal_root = Path(client.app.state.extension_runtime.proposal_root)
    original_mkdir = Path.mkdir

    def broken_mkdir(self, *args, **kwargs):
        if self.parent == proposal_root:  # 只拦候选目录，不影响测试基础设施
            raise OSError("cannot create candidate directory")
        return original_mkdir(self, *args, **kwargs)

    monkeypatch.setattr(Path, "mkdir", broken_mkdir)

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 503
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "ARTIFACT_WRITE_FAILED"
    assert body["message_key"] == "errors.extension.artifactWriteFailed"
    # 失败发生在授权消费之前：授权保留，可刷新后重试；无残留、无 Artifact
    assert grant_store(client).active_count() == 1
    assert candidate_files(client) == []
    with client.app.state.service.database.sessions() as session:
        assert session.execute(text("SELECT COUNT(*) FROM artifacts")).scalar_one() == 0


def test_candidate_save_os_error_is_contract_artifact_write_failed(
    tmp_path, tiny_workspace, monkeypatch
):
    """workbook.save 的 OS 级失败（磁盘满/AV 锁定）契约化，绝不逃逸为 500。"""
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)

    def broken_save(self, path):
        raise OSError("disk full")

    monkeypatch.setattr(
        "dst_manager.extensions.builtin.sheet_catalog.workbook.Workbook.save", broken_save
    )

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 503
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "ARTIFACT_WRITE_FAILED"
    assert body["message_key"] == "errors.extension.artifactWriteFailed"
    # 失败发生在授权消费之前：授权保留；目标未写入、无残留、无 Artifact
    assert grant_store(client).active_count() == 1
    assert not (exports / "out.xlsx").exists()
    assert candidate_files(client) == []
    with client.app.state.service.database.sessions() as session:
        assert session.execute(text("SELECT COUNT(*) FROM artifacts")).scalar_one() == 0


def test_replace_failure_keeps_old_target_without_artifact(tmp_path, tiny_workspace, monkeypatch):
    """批次三检查点：os.replace 失败 → 旧目标字节保持、零半文件、零 Artifact。"""
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    target = exports / "out.xlsx"
    target.write_bytes(b"old user file")
    grant = create_grant(grant_store(client), workspace, target)
    preview_body = preview_action(client, workspace)

    def broken_replace(src, dst):
        raise OSError("replace locked")

    monkeypatch.setattr("dst_manager.extensions.artifacts.os.replace", broken_replace)

    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)

    assert resp.status_code == 503
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "ARTIFACT_WRITE_FAILED"
    # 旧目标字节保持、目标目录无半文件残留、无 Artifact 登记
    assert target.read_bytes() == b"old user file"
    assert list(exports.iterdir()) == [target]
    assert candidate_files(client) == []
    with client.app.state.service.database.sessions() as session:
        assert session.execute(text("SELECT COUNT(*) FROM artifacts")).scalar_one() == 0


def test_execute_without_grant_channel_returns_contract_503(tmp_path, tiny_workspace):
    """授权通道未装配（如最小测试装配）时契约化拒绝，绝不 AttributeError。"""
    settings = Settings(data_dir=tmp_path / "data")
    settings.data_dir.mkdir(parents=True)
    database = Database(settings.database_url)
    runtime = ExtensionRuntime(
        ExtensionRegistry(),
        ExtensionStore(database.sessions),
        reader=ExtensionWorkspaceReader(database.sessions),
    )
    client = TestClient(create_app(settings, extension_runtime=runtime))
    workspace = open_workspace(client, tiny_workspace)
    preview_body = preview_action(client, workspace)

    resp = execute_action(client, workspace, preview_body, "any-grant")

    assert resp.status_code == 503
    body = resp.json()
    assert_error_contract(body)
    assert body["code"] == "EXTENSION_CAPABILITY_UNAVAILABLE"


# ---------------------------------------------------------------------------
# Artifact 三态可用性派生
# ---------------------------------------------------------------------------


def test_artifact_availability_tracks_filesystem_state(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    preview_body = preview_action(client, workspace)
    executed = execute_action(client, workspace, preview_body, grant.save_grant_id)
    artifact_id = executed.json()["artifact_id"]
    output = Path(executed.json()["output_path"])

    # 文件被删除 → MISSING（历史记录不伪装成当前可用）
    output.unlink()
    assert client.get(f"/api/artifacts/{artifact_id}").json()["availability"] == "MISSING"

    # 文件被替换为不同内容 → CHANGED
    output.write_bytes(b"different content")
    assert client.get(f"/api/artifacts/{artifact_id}").json()["availability"] == "CHANGED"


# ---------------------------------------------------------------------------
# 桌面共同装配：桥与 API runtime 共享同一个 SaveGrantStore
# ---------------------------------------------------------------------------


class _FakeWindow:
    def __init__(self, result):
        self._result = result

    def create_file_dialog(self, *args, **kwargs):
        return self._result


class _FakeContext:
    """可信上下文替身：只暴露桥消费的最小面。"""

    def __init__(self, workspace_id: str, dst_path: Path):
        self._workspace_id = workspace_id
        self.dst_path = dst_path
        self.root = dst_path.parent

    @property
    def current(self):
        return self

    @property
    def workspace_id(self) -> str:
        return self._workspace_id


def test_bridge_created_grant_is_consumable_by_api_execute(tmp_path, tiny_workspace):
    """run_desktop 共同装配的行为钉子：桥创建的授权必须能被 API 执行消费。"""
    store = SaveGrantStore()
    client = make_client(tmp_path, save_grants=store)
    assert client.app.state.extension_runtime.save_grants is store
    workspace = open_workspace(client, tiny_workspace)
    exports = tmp_path / "exports"
    exports.mkdir()
    bridge = ShellBridge(
        context=_FakeContext(workspace["id"], tiny_workspace[0]),
        registry=client.app.state.extension_runtime.registry,
        save_grants=store,
    )
    bridge.bind(_FakeWindow(str(exports / "桥选.xlsx")))

    result = bridge.request_extension_save(SHEET_CATALOG_ID, ACTION_ID, workspace["id"])

    assert result["ok"] is True
    grant_id = result["value"]["save_grant_id"]
    preview_body = preview_action(client, workspace)
    resp = execute_action(client, workspace, preview_body, grant_id)
    assert resp.status_code == 200, resp.text
    assert Path(resp.json()["output_path"]).name == "桥选.xlsx"


# ---------------------------------------------------------------------------
# 只读不变量：预览→执行前后工程与数据库零写入
# ---------------------------------------------------------------------------


def test_preview_and_execute_leave_project_and_database_read_only(tmp_path, tiny_workspace):
    client = make_client(tmp_path, save_grants=SaveGrantStore())
    workspace = open_workspace(client, tiny_workspace)
    dst_path = tiny_workspace[0]
    dwg_path = dst_path.parent / "A.dwg"
    project_root = dst_path.parent
    data_dir = tmp_path / "data"
    exports = tmp_path / "exports"
    exports.mkdir()

    def project_tree():
        # 测试布局折衷：应用数据目录与用户选择的保存目标都位于工程目录内，
        # 二者都不属于工程内容，排除后只对工程文件树做不变量比对。
        return {
            str(p.relative_to(project_root)): (p.stat().st_size, p.stat().st_mtime_ns)
            for p in project_root.rglob("*")
            if p.is_file() and data_dir not in p.parents and exports not in p.parents
        }

    def db_counts():
        with client.app.state.service.database.sessions() as session:
            return {
                table: session.execute(text(f"SELECT COUNT(*) FROM {table}")).scalar_one()
                for table in ("jobs", "document_revisions", "workspace_write_locks")
            }

    before_tree = project_tree()
    before_counts = db_counts()
    before_bytes = {p: p.read_bytes() for p in (dst_path, dwg_path)}
    before_mtime = {p: p.stat().st_mtime_ns for p in (dst_path, dwg_path)}

    preview_body = preview_action(client, workspace)
    assert preview_body["executable"] is True
    grant = create_grant(grant_store(client), workspace, exports / "out.xlsx")
    resp = execute_action(client, workspace, preview_body, grant.save_grant_id)
    assert resp.status_code == 200

    assert project_tree() == before_tree
    assert db_counts() == before_counts
    for path in (dst_path, dwg_path):
        assert path.read_bytes() == before_bytes[path]
        assert path.stat().st_mtime_ns == before_mtime[path]
    assert candidate_files(client) == []
