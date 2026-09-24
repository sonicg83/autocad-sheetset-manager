"""创建草稿、XLSX 与权威预览 API 契约（PLAN-DM-036 Task 4 / SPEC-DM-018 §3.2、§5、§6）。

覆盖：草稿 CRUD 与乐观修订；标准候选只列已发布且依赖、受控资产可用的标准
（缺依赖/缺模板文件的候选不可选且带原因）；XLSX 模板导出与全量原子导入
（非法导入零变更、陈旧修订 409、成功导入整批替换）；按组预览表与逐张属性
明细；`preview_digest` 对标准文档、草稿、有效设置、资产与目标状态的敏感性；
非空目标与空图纸组阻断；预览全程不启动 CAD。

Task 6 追加：`/{id}/execute` 只接受草稿 ID 与 `preview_digest`，执行前重新加载
标准、草稿、目标与设置/资产快照并重算摘要（漂移 409 `CREATION_PREVIEW_STALE`），
入队创建任务且不写目标目录。

测试只用 `tmp_path` 夹具，不依赖用户本地数据库或上次运行残留。
"""

import copy
import shutil
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from yaml import safe_dump

from dst_manager.application.creation_import import build_creation_template
from dst_manager.config import Settings
from dst_manager.domain.creation import CreationAssetOption
from dst_manager.domain.standards import parse_published_standard_document
from dst_manager.extensions.builtin.index import BuiltinExtensionEntry
from dst_manager.infrastructure.autocad.worker import CoreConsoleExecutor
from dst_manager.infrastructure.creation_xlsx import (
    LIST_SHEET,
    META_SHEET,
    SHEET_SHEET,
    SHEETSET_SHEET,
    TARGET_PATH_LABEL,
)
from dst_manager.interfaces.api import create_app

#: 最小可发布标准：两个普通 sheetset 属性、一个派生映射、一个普通 sheet 文本、
#: 一个普通 sheet 枚举、一个派生组合，外加基础模板与布局模板各一个受控资产。
STANDARD_DOCUMENT: dict[str, object] = {
    "schema_version": 1,
    "standard_id": "szmedi.gas",
    "version": "2.1.0",
    "name": "市政燃气施工图",
    "supported_cad_versions": ["2016", "2020"],
    "properties": [
        {
            "property_id": "prop-name",
            "name": "工程名称",
            "scope": "sheetset",
            "kind": "text",
            "required": True,
            "default_value": "默认工程",
        },
        {
            "property_id": "prop-major",
            "name": "专业",
            "scope": "sheetset",
            "kind": "enum",
            "required": True,
            "default_value": "燃气",
            "enum_items": [
                {"item_id": "enum-gas", "value": "燃气"},
                {"item_id": "enum-oil", "value": "燃油"},
            ],
        },
        {
            "property_id": "prop-code",
            "name": "专业代码",
            "scope": "sheetset",
            "kind": "mapping",
            "source_property_id": "prop-major",
            "mapping": [
                {"item_id": "enum-gas", "value": "RQ"},
                {"item_id": "enum-oil", "value": "RY"},
            ],
            "confirmed_source_items": [["enum-gas", "燃气"], ["enum-oil", "燃油"]],
        },
        {
            "property_id": "prop-stage",
            "name": "设计阶段",
            "scope": "sheet",
            "kind": "text",
            "default_value": "施工图",
        },
        {
            "property_id": "prop-part",
            "name": "分部",
            "scope": "sheet",
            "kind": "enum",
            "enum_items": [
                {"item_id": "enum-part-a", "value": "A 段"},
                {"item_id": "enum-part-b", "value": "B 段"},
            ],
        },
        {
            "property_id": "prop-label",
            "name": "图签",
            "scope": "sheet",
            "kind": "composition",
            "segments": [
                {"property_id": "prop-code"},
                {"literal": "-"},
                {"system_field": "sheet.number"},
            ],
        },
    ],
    "dwg_naming": {
        "segments": [
            {"property_id": "prop-code"},
            {"literal": "-"},
            {"system_field": "subset.scope"},
            {"literal": " "},
            {"system_field": "subset.name"},
        ]
    },
    "assets": [
        {
            "asset_id": "base-a1",
            "kind": "base-template",
            "files": [{"path": "templates/a1.dwt", "role": ""}],
        },
        {
            "asset_id": "layout-a1",
            "kind": "layout-template",
            "files": [{"path": "templates/a1-layout.dwt", "role": "A1"}],
        },
    ],
    "numbering": {"sequence_field": "subset.sequence", "digits": 2},
}

#: 声明资产文件缺失的标准：同一文档换身份，模板文件不写盘。
MISSING_ASSET_DOCUMENT: dict[str, object] = {
    **STANDARD_DOCUMENT,
    "standard_id": "szmedi.missing",
    "version": "1.0.0",
}

#: 同类内文件名冲突的标准：两个基础模板同名不同目录，标签必须退回包内相对路径。
DUPLICATE_LABEL_DOCUMENT: dict[str, object] = {
    **STANDARD_DOCUMENT,
    "standard_id": "szmedi.dupe",
    "version": "1.0.0",
    "assets": [
        {"asset_id": "base-a", "kind": "base-template", "files": [{"path": "a1/a1.dwt"}]},
        {"asset_id": "base-b", "kind": "base-template", "files": [{"path": "b1/a1.dwt"}]},
        {
            "asset_id": "layout-a1",
            "kind": "layout-template",
            "files": [{"path": "templates/a1-layout.dwt", "role": "A1"}],
        },
    ],
}

DUPLICATE_LABEL_FILES: dict[str, bytes] = {
    "a1/a1.dwt": b"base-a",
    "b1/a1.dwt": b"base-b",
    "templates/a1-layout.dwt": b"layout",
}

#: 声明受信扩展依赖的标准（发布时该扩展必须已注册）。
DEPENDENT_DOCUMENT: dict[str, object] = {
    **STANDARD_DOCUMENT,
    "standard_id": "szmedi.dependent",
    "version": "1.0.0",
    "dependencies": [
        {
            "extension_id": "test.trusted",
            "capability_id": "layout.template.v1",
            "min_version": "1.0.0",
        },
    ],
}

#: 受控资产文件内容：预览只哈希内容，不读 DWG，也不启动 CAD。
ASSET_FILES: dict[str, bytes] = {
    "templates/a1.dwt": b"base-template-bytes",
    "templates/a1-layout.dwt": b"layout-template-bytes",
}

#: 与发布标准一致（标签取包内受控文件名）的资产候选，供本地构造模板使用。
ASSET_OPTIONS: tuple[CreationAssetOption, ...] = (
    CreationAssetOption(asset_id="base-a1", kind="base-template", label="a1.dwt"),
    CreationAssetOption(
        asset_id="layout-a1", kind="layout-template", label="a1-layout.dwt", layouts=("A1",)
    ),
)

#: 可输入 sheetset 普通属性（派生属性不是输入项）。
SHEETSET_VALUES: dict[str, str] = {"prop-name": "示例工程", "prop-major": "燃气"}
#: 可输入 sheet 普通属性：同组全部图纸继承同一份输入。
SHEET_VALUES: dict[str, str] = {"prop-stage": "施工图", "prop-part": "A 段"}


def group_input(**overrides: object) -> dict[str, object]:
    """一个完整图纸组输入（字段与草稿保存契约一致）。"""
    group: dict[str, object] = {
        "group_id": "group-1",
        "created_order": 1,
        "title": "平面图",
        "count": 2,
        "base_asset_id": "base-a1",
        "layout_asset_id": "layout-a1",
        "paper_layout": "A1",
        "sheet_values": dict(SHEET_VALUES),
    }
    group.update(overrides)
    return group


@dataclass(frozen=True, slots=True)
class DraftFixture:
    """创建草稿夹具：只暴露测试需要的草稿身份。"""

    id: str


@dataclass(frozen=True, slots=True)
class PreviewFixture:
    """权威预览夹具：只暴露执行入口需要的摘要。"""

    digest: str


def make_client(tmp_path: Path, manifest: dict | None = None) -> TestClient:
    """测试客户端；``manifest`` 用于注册受信扩展（校验标准依赖）。"""
    kwargs: dict[str, object] = {}
    if manifest is not None:
        manifest_path = tmp_path / "trusted-extension.yaml"
        manifest_path.write_text(safe_dump(manifest), encoding="utf-8")
        kwargs["extension_index"] = [
            BuiltinExtensionEntry(manifest_resource=str(manifest_path), factory=_NoopExtension)
        ]
    return TestClient(create_app(Settings(data_dir=tmp_path / "data"), **kwargs))


class _NoopExtension:
    def start(self) -> None:
        return None

    def stop(self) -> None:
        return None


TRUSTED_MANIFEST: dict[str, object] = {
    "extension_id": "test.trusted",
    "version": "1.2.0",
    "extension_type": "builtin",
    "host_contract": 1,
    "enabled_by_default": True,
    "name_key": "extensions.trusted.name",
    "description_key": "extensions.trusted.description",
    "required_capabilities": [],
    "provided_capabilities": [{"capability_id": "layout.template.v1", "version": "1.0.0"}],
    "actions": [],
    "settings_schema": 1,
}


def data_dir(tmp_path: Path) -> Path:
    return tmp_path / "data"


def publish_standard(
    client: TestClient,
    root: Path,
    *,
    document: dict | None = None,
    asset_files: dict[str, bytes] | None = ASSET_FILES,
    draft_id: str = "draft-gas",
) -> None:
    """发布标准：资产文件先写入草稿目录，发布后随目录一起落到 published 下。"""
    response = client.post(
        "/api/standards/drafts",
        json={"draft_id": draft_id, "document": document or STANDARD_DOCUMENT},
    )
    assert response.status_code == 200, response.text
    for relative, content in (asset_files or {}).items():
        target = root / "standards" / "user" / "drafts" / draft_id / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(content)
    published = client.post(f"/api/standards/drafts/{draft_id}/publish")
    assert published.status_code == 200, published.text


def published_root(root: Path, standard_id: str = "szmedi.gas", version: str = "2.1.0") -> Path:
    return root / "standards" / "user" / "published" / standard_id / version


def create_creation_draft(client: TestClient) -> dict:
    response = client.post(
        "/api/creation-drafts", json={"standard_id": "szmedi.gas", "version": "2.1.0"}
    )
    assert response.status_code == 200, response.text
    return response.json()


def save_creation_draft(
    client: TestClient,
    draft_id: str,
    *,
    expected_revision: int,
    target_path: str,
    groups: list[dict] | None = None,
    sheetset_values: dict[str, str] | None = None,
    step: str = "review",
) -> dict:
    response = client.put(
        f"/api/creation-drafts/{draft_id}",
        json={
            "expected_revision": expected_revision,
            "step": step,
            "target_path": target_path,
            "sheetset_values": sheetset_values or dict(SHEETSET_VALUES),
            "groups": [group_input()] if groups is None else groups,
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def update_creation_draft(client: TestClient, draft_id: str) -> dict:
    """编辑草稿（改一个图纸集输入值）并返回保存后的草稿。"""
    current = client.get(f"/api/creation-drafts/{draft_id}").json()
    return save_creation_draft(
        client,
        draft_id,
        expected_revision=current["revision"],
        target_path=current["target_path"],
        groups=current["groups"],
        sheetset_values={**current["sheetset_values"], "prop-name": "编辑后工程"},
        step=current["step"],
    )


def preview_draft(client: TestClient, draft_id: str):
    return client.post(f"/api/creation-drafts/{draft_id}/preview")


def fill_template(
    data: bytes, sheetset: dict[str, str], rows: tuple[dict[str, object], ...]
) -> bytes:
    """按可见属性名/表头填写模板：``SheetSet`` 按 A 列标签定位，``Sheet`` 按表头定位。"""
    workbook = load_workbook(BytesIO(data))
    sheet_set = workbook[SHEETSET_SHEET]
    for label, value in sheetset.items():
        for row in range(1, sheet_set.max_row + 1):
            if sheet_set.cell(row=row, column=1).value == label:
                sheet_set.cell(row=row, column=2).value = value
                break
        else:  # pragma: no cover - 夹具模板必然含该行标签
            raise AssertionError(f"模板缺少行标签 {label!r}")
    sheet = workbook[SHEET_SHEET]
    headers = {
        sheet.cell(row=1, column=column).value: column
        for column in range(1, sheet.max_column + 1)
    }
    for index, values in enumerate(rows, start=2):
        for header, value in values.items():
            column = headers.get(header)
            assert column is not None, f"模板缺少表头 {header!r}"
            sheet.cell(row=index, column=column).value = value
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def valid_xlsx(target_path: str) -> bytes:
    """本地构造的合法工作簿：与发布标准的资产候选、属性列一一对应。"""
    template = build_creation_template(
        parse_published_standard_document(STANDARD_DOCUMENT), ASSET_OPTIONS
    )
    return fill_template(
        template,
        {TARGET_PATH_LABEL: target_path, "工程名称": "导入工程", "专业": "燃油"},
        (
            {
                "图名": "平面图",
                "张数": "2",
                "基础模板": "a1.dwt",
                "布局模板": "a1-layout.dwt",
                "图幅": "A1",
                "设计阶段": "施工图",
                "分部": "B 段",
            },
        ),
    )


def forbid_cad(monkeypatch: pytest.MonkeyPatch) -> None:
    """把 CAD 读取与 Core Console 执行都替换成失败：预览一旦触达即测试失败。"""

    def fail(*args: object, **kwargs: object):
        raise AssertionError("预览全过程不得启动 CAD")

    monkeypatch.setattr(CoreConsoleExecutor, "run", fail)
    monkeypatch.setattr(
        "dst_manager.application.service.DstManagerService.get_layout_names", fail
    )


@pytest.fixture
def client(tmp_path: Path) -> TestClient:
    return make_client(tmp_path)


@pytest.fixture
def root(tmp_path: Path) -> Path:
    return data_dir(tmp_path)


@pytest.fixture
def creation_draft(client: TestClient, root: Path, tmp_path: Path) -> DraftFixture:
    """发布标准并保存一份完整可执行的创建草稿（目标为尚不存在的新目录）。"""
    publish_standard(client, root)
    draft = create_creation_draft(client)
    saved = save_creation_draft(
        client,
        draft["id"],
        expected_revision=draft["revision"],
        target_path=str(tmp_path / "projects" / "新建项目"),
    )
    assert saved["groups"] == [group_input()]
    return DraftFixture(id=draft["id"])


@pytest.fixture
def preview(client: TestClient, creation_draft: DraftFixture) -> PreviewFixture:
    """权威预览夹具：执行入口只接受这里算出的 ``preview_digest``。"""
    return PreviewFixture(
        digest=preview_draft(client, creation_draft.id).json()["preview_digest"]
    )


@pytest.fixture
def nonempty_target(client: TestClient, creation_draft: DraftFixture, tmp_path: Path) -> Path:
    """把草稿目标改成一个已存在且非空的目录（模拟预览前被其他程序占用）。"""
    target = tmp_path / "occupied" / "新建项目"
    target.mkdir(parents=True)
    (target / "已有工程.dst").write_text("x", encoding="utf-8")
    current = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    save_creation_draft(
        client,
        creation_draft.id,
        expected_revision=current["revision"],
        target_path=str(target),
        groups=current["groups"],
        sheetset_values=current["sheetset_values"],
    )
    return target


@pytest.fixture
def invalid_xlsx() -> bytes:
    """可读但结构非法的模板：``SheetSet`` 表头被改坏，导入必须整批拒绝。"""
    template = build_creation_template(
        parse_published_standard_document(STANDARD_DOCUMENT), ASSET_OPTIONS
    )
    workbook = load_workbook(BytesIO(template))
    workbook[SHEETSET_SHEET].cell(row=1, column=1).value = "字段名"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---- Step 1：三个强制用例 --------------------------------------------------


def test_preview_digest_changes_after_draft_edit(client: TestClient, creation_draft: DraftFixture) -> None:
    preview = client.post(f"/api/creation-drafts/{creation_draft.id}/preview").json()
    update_creation_draft(client, creation_draft.id)
    updated = client.post(f"/api/creation-drafts/{creation_draft.id}/preview").json()
    assert updated["preview_digest"] != preview["preview_digest"]


def test_invalid_import_keeps_draft_unchanged(
    client: TestClient, creation_draft: DraftFixture, invalid_xlsx: bytes
) -> None:
    before = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    response = client.post(
        f"/api/creation-drafts/{creation_draft.id}/xlsx-import", files={"file": invalid_xlsx}
    )
    assert response.status_code == 422
    assert client.get(f"/api/creation-drafts/{creation_draft.id}").json() == before


def test_nonempty_target_blocks_preview(
    client: TestClient, creation_draft: DraftFixture, nonempty_target: Path
) -> None:
    response = client.post(f"/api/creation-drafts/{creation_draft.id}/preview")
    assert response.status_code == 422
    assert response.json()["code"] == "CREATION_TARGET_NOT_EMPTY"


def test_execute_rejects_stale_preview(
    client: TestClient, creation_draft: DraftFixture, preview: PreviewFixture
) -> None:
    update_creation_draft(client, creation_draft.id)
    response = client.post(
        f"/api/creation-drafts/{creation_draft.id}/execute",
        json={"preview_digest": preview.digest},
    )
    assert response.status_code == 409
    assert response.json()["code"] == "CREATION_PREVIEW_STALE"


def test_execute_rejects_non_executable_draft(
    client: TestClient, root: Path, tmp_path: Path
) -> None:
    """阻断诊断（空图纸组）不得入队：执行入口以稳定码拒绝。"""
    publish_standard(client, root)
    draft = create_creation_draft(client)
    save_creation_draft(
        client,
        draft["id"],
        expected_revision=draft["revision"],
        target_path=str(tmp_path / "projects" / "新建项目"),
        groups=[],
    )
    body = preview_draft(client, draft["id"]).json()
    assert body["executable"] is False

    response = client.post(
        f"/api/creation-drafts/{draft['id']}/execute",
        json={"preview_digest": body["preview_digest"]},
    )

    assert response.status_code == 422
    assert response.json()["code"] == "CREATION_PLAN_INVALID"


def test_execute_enqueues_creation_job_without_workspace(
    client: TestClient, creation_draft: DraftFixture, preview: PreviewFixture
) -> None:
    """执行只入队：任务没有普通工作区，也不写目标目录（成果只能由 Worker 发布）。"""
    body = preview_draft(client, creation_draft.id).json()
    target = Path(body["target_path"])

    response = client.post(
        f"/api/creation-drafts/{creation_draft.id}/execute",
        json={"preview_digest": preview.digest},
    )

    assert response.status_code == 200, response.text
    job = response.json()
    assert job["type"] == "creation"
    assert job["status"] == "QUEUED"
    assert job["workspace_id"] is None
    assert job["creation_draft_id"] == creation_draft.id
    assert job["payload"]["creation_draft_id"] == creation_draft.id
    assert job["payload"]["preview_digest"] == preview.digest
    assert job["payload"]["target_path"] == body["target_path"]
    assert job["payload"]["standard"] == {"standard_id": "szmedi.gas", "version": "2.1.0"}
    assert not target.exists()
    # 无工作区时仍能记录状态与时间线（SSE 事件同源）
    detail = client.get(f"/api/jobs/{job['id']}").json()
    assert detail["workspace_id"] is None
    assert [item["status"] for item in detail["timeline"]] == ["QUEUED"]
    assert detail["files"] == []


def test_execute_keeps_draft_unchanged_and_does_not_start_cad(
    client: TestClient, creation_draft: DraftFixture, preview: PreviewFixture, monkeypatch
) -> None:
    forbid_cad(monkeypatch)
    before = client.get(f"/api/creation-drafts/{creation_draft.id}").json()

    assert client.post(
        f"/api/creation-drafts/{creation_draft.id}/execute",
        json={"preview_digest": preview.digest},
    ).status_code == 200

    assert client.get(f"/api/creation-drafts/{creation_draft.id}").json() == before


def test_job_events_stream_stops_at_needs_review(
    client: TestClient, creation_draft: DraftFixture, preview: PreviewFixture
) -> None:
    """创建任务可终局于 NEEDS_REVIEW：SSE 必须在该状态收尾（不能继续轮询）。

    创建发布日志身份不可证明时任务被隔离为 `NEEDS_REVIEW`（Task 6/7）。事件流的终止
    集合曾漏掉该状态，任务已终局而流不结束，客户端只能靠断线回退才收到终态。
    """
    job = client.post(
        f"/api/creation-drafts/{creation_draft.id}/execute",
        json={"preview_digest": preview.digest},
    ).json()
    database = client.app.state.service.database
    assert database.finalize_job_terminal(
        job["id"], "NEEDS_REVIEW", "CREATION_PUBLISH_REVIEW_REQUIRED", "创建发布日志身份不可证明"
    )

    # 流必须自己结束（否则 iter_text 会一直等待），且推送的最后一帧就是终态
    with client.stream("GET", f"/api/jobs/{job['id']}/events") as response:
        body = "".join(response.iter_text())

    assert "NEEDS_REVIEW" in body
    assert client.get(f"/api/jobs/{job['id']}").json()["status"] == "NEEDS_REVIEW"


# ---- 草稿 CRUD ------------------------------------------------------------


def test_draft_crud_roundtrip(client: TestClient, creation_draft: DraftFixture) -> None:
    draft = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    assert draft["standard_id"] == "szmedi.gas"
    assert draft["standard_version"] == "2.1.0"
    assert draft["revision"] == 2

    saved = update_creation_draft(client, creation_draft.id)
    assert saved["revision"] == 3
    assert saved["sheetset_values"]["prop-name"] == "编辑后工程"

    stale = client.put(
        f"/api/creation-drafts/{creation_draft.id}",
        json={
            "expected_revision": 1,
            "step": "review",
            "target_path": draft["target_path"],
            "sheetset_values": draft["sheetset_values"],
            "groups": draft["groups"],
        },
    )
    assert stale.status_code == 409
    assert stale.json()["code"] == "CREATION_DRAFT_CONFLICT"

    assert client.delete(f"/api/creation-drafts/{creation_draft.id}").status_code == 200
    assert client.get(f"/api/creation-drafts/{creation_draft.id}").status_code == 404


def test_unknown_draft_is_stable_not_found(client: TestClient) -> None:
    assert client.get("/api/creation-drafts/absent").status_code == 404
    assert preview_draft(client, "absent").status_code == 404


def test_draft_creation_requires_published_standard(client: TestClient) -> None:
    response = client.post(
        "/api/creation-drafts", json={"standard_id": "szmedi.gas", "version": "2.1.0"}
    )
    assert response.status_code == 404
    assert response.json()["code"] == "CREATION_STANDARD_MISSING"


def test_execute_endpoint_requires_preview_digest(
    client: TestClient, creation_draft: DraftFixture
) -> None:
    """执行端点只接受草稿 ID 与 preview_digest：缺失/多余派生输出都被契约拒绝。"""
    paths = [getattr(route, "path", "") for route in client.app.routes]
    assert "/api/creation-drafts/{draft_id}/execute" in paths
    missing = client.post(f"/api/creation-drafts/{creation_draft.id}/execute", json={})
    assert missing.status_code == 422
    derived = client.post(
        f"/api/creation-drafts/{creation_draft.id}/execute",
        json={"preview_digest": "x", "target_path": "C:\\other"},
    )
    assert derived.status_code == 422


# ---- 标准候选 -------------------------------------------------------------


def test_candidates_list_available_standard_with_asset_options(
    client: TestClient, root: Path
) -> None:
    publish_standard(client, root)
    candidates = client.get("/api/creation-drafts/standards").json()
    assert candidates == [
        {
            "standard_id": "szmedi.gas",
            "version": "2.1.0",
            "name": "市政燃气施工图",
            "supported_cad_versions": ["2016", "2020"],
            "available": True,
            "reasons": [],
            "asset_options": [
                {"asset_id": "base-a1", "kind": "base-template", "label": "a1.dwt", "layouts": []},
                {
                    "asset_id": "layout-a1",
                    "kind": "layout-template",
                    "label": "a1-layout.dwt",
                    "layouts": ["A1"],
                },
            ],
        }
    ]


def test_candidate_without_template_files_is_unavailable(
    client: TestClient, root: Path
) -> None:
    publish_standard(client, root)
    publish_standard(
        client, root, document=MISSING_ASSET_DOCUMENT, draft_id="draft-missing"
    )
    # 发布门禁保证资产在发布时必须存在；这里模拟发布后模板文件被外部删除/移动，
    # 候选列表仍必须给出稳定原因而不是把坏标准当成可选。
    shutil.rmtree(
        published_root(root, standard_id="szmedi.missing", version="1.0.0") / "templates"
    )
    candidates = {
        item["standard_id"]: item
        for item in client.get("/api/creation-drafts/standards").json()
    }
    assert candidates["szmedi.gas"]["available"] is True
    broken = candidates["szmedi.missing"]
    assert broken["available"] is False
    assert broken["reasons"] == [
        "标准资产 'base-a1' 声明的文件 'templates/a1.dwt' 不存在或路径非法",
        "标准资产 'layout-a1' 声明的文件 'templates/a1-layout.dwt' 不存在或路径非法",
    ]


@pytest.mark.parametrize(
    ("case", "content"),
    [
        ("truncated", '{"schema_version": 1'),
        ("not-utf8", None),
    ],
)
def test_candidate_with_corrupt_standard_document_is_unavailable(
    client: TestClient, root: Path, case: str, content: str | None
) -> None:
    """已发布但内容损坏的 document.json：候选列表必须仍返回 200 并给出稳定原因。"""
    publish_standard(client, root)
    document = published_root(root) / "document.json"
    if content is None:
        document.write_bytes(b"\xff\xfe\x00\x01")
    else:
        document.write_text(content, encoding="utf-8")

    response = client.get("/api/creation-drafts/standards")

    assert response.status_code == 200, case
    candidates = {item["standard_id"]: item for item in response.json()}
    broken = candidates["szmedi.gas"]
    assert broken["available"] is False
    assert len(broken["reasons"]) == 1 and "标准文档无法解析" in broken["reasons"][0]


def test_candidates_survive_illegal_published_directory_name(
    client: TestClient, root: Path
) -> None:
    """标准库中混入非法版本目录名时，候选列表仍返回 200 并列出其余标准。"""
    publish_standard(client, root)
    (root / "standards" / "user" / "published" / "szmedi.gas" / "tmp").mkdir()

    response = client.get("/api/creation-drafts/standards")

    assert response.status_code == 200, response.text
    assert [item["standard_id"] for item in response.json()] == ["szmedi.gas"]


def test_candidate_labels_are_unique_within_kind(client: TestClient, root: Path) -> None:
    """同类内文件名冲突时标签退回包内相对路径：候选标签必须可唯一回指资产。"""
    publish_standard(
        client,
        root,
        document=DUPLICATE_LABEL_DOCUMENT,
        asset_files=DUPLICATE_LABEL_FILES,
        draft_id="draft-dupe",
    )
    candidates = {
        item["standard_id"]: item
        for item in client.get("/api/creation-drafts/standards").json()
    }
    assert candidates["szmedi.dupe"]["available"] is True
    assert [item["label"] for item in candidates["szmedi.dupe"]["asset_options"]] == [
        "a1/a1.dwt",
        "b1/a1.dwt",
        "a1-layout.dwt",
    ]


def test_candidate_with_missing_trusted_dependency_is_unavailable(
    tmp_path: Path, root: Path
) -> None:
    """依赖可用的注册表里发布，随后在没有该扩展的进程里候选必须不可选。"""
    with_extension = make_client(tmp_path, TRUSTED_MANIFEST)
    publish_standard(with_extension, root, document=DEPENDENT_DOCUMENT, draft_id="draft-dep")
    without_extension = make_client(tmp_path)
    candidates = {
        item["standard_id"]: item
        for item in without_extension.get("/api/creation-drafts/standards").json()
    }
    dependent = candidates["szmedi.dependent"]
    assert dependent["available"] is False
    assert dependent["reasons"] == [
        "受信扩展依赖未满足：test.trusted/layout.template.v1（extension-missing）"
    ]


# ---- XLSX 模板与导入 ------------------------------------------------------


def test_xlsx_template_exports_current_standard(
    client: TestClient, creation_draft: DraftFixture
) -> None:
    response = client.get(f"/api/creation-drafts/{creation_draft.id}/xlsx-template")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "creation-template-szmedi.gas-2.1.0.xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == [SHEETSET_SHEET, SHEET_SHEET, META_SHEET, LIST_SHEET]


def test_import_replaces_all_inputs_atomically(
    client: TestClient, creation_draft: DraftFixture, tmp_path: Path
) -> None:
    before = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    target = tmp_path / "imported" / "新建项目"
    response = client.post(
        f"/api/creation-drafts/{creation_draft.id}/xlsx-import",
        files={"file": valid_xlsx(str(target))},
        params={"expected_revision": before["revision"]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["revision"] == before["revision"] + 1
    assert body["target_path"] == str(target)
    assert body["sheetset_values"] == {"prop-name": "导入工程", "prop-major": "燃油"}
    assert body["groups"] == [
        group_input(
            group_id="xlsx-2",
            created_order=1,
            title="平面图",
            count=2,
            sheet_values={"prop-stage": "施工图", "prop-part": "B 段"},
        )
    ]
    # 导入是覆盖式的：旧路径与旧组被整批替换，不留半新半旧状态。
    assert body["target_path"] != before["target_path"]


def test_import_with_stale_revision_keeps_draft_unchanged(
    client: TestClient, creation_draft: DraftFixture, tmp_path: Path
) -> None:
    before = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    response = client.post(
        f"/api/creation-drafts/{creation_draft.id}/xlsx-import",
        files={"file": valid_xlsx(str(tmp_path / "imported" / "新建项目"))},
        params={"expected_revision": before["revision"] - 1},
    )
    assert response.status_code == 409
    assert response.json()["code"] == "CREATION_DRAFT_CONFLICT"
    assert client.get(f"/api/creation-drafts/{creation_draft.id}").json() == before


def test_import_rejection_reports_locatable_diagnostics(
    client: TestClient, creation_draft: DraftFixture, invalid_xlsx: bytes
) -> None:
    response = client.post(
        f"/api/creation-drafts/{creation_draft.id}/xlsx-import", files={"file": invalid_xlsx}
    )
    assert response.status_code == 422
    body = response.json()
    assert body["code"] == "CREATION_IMPORT_INVALID"
    assert body["message_key"] == "errors.creation.importInvalid"
    assert body["diagnostics"] == [
        {
            "code": "CREATION_XLSX_HEADER_MISMATCH",
            "message": "SheetSet 表头 '字段名' 与模板 '属性名' 不一致",
            "sheet": SHEETSET_SHEET,
            "row": 1,
            "column": "A",
        }
    ]


def test_import_rejects_oversized_workbook(
    client: TestClient, creation_draft: DraftFixture, monkeypatch: pytest.MonkeyPatch
) -> None:
    """超出受控大小的输入按受控规模拒绝，不进解析、不改草稿。"""
    monkeypatch.setattr("dst_manager.application.creation.MAX_CREATION_XLSX_BYTES", 8)
    before = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    response = client.post(
        f"/api/creation-drafts/{creation_draft.id}/xlsx-import",
        files={"file": b"0123456789"},
    )
    assert response.status_code == 422
    assert response.json()["diagnostics"][0]["code"] == "CREATION_XLSX_SCALE_EXCEEDED"
    assert client.get(f"/api/creation-drafts/{creation_draft.id}").json() == before


def test_import_rejects_unreadable_workbook(
    client: TestClient, creation_draft: DraftFixture
) -> None:
    before = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    response = client.post(
        f"/api/creation-drafts/{creation_draft.id}/xlsx-import",
        files={"file": b"not-a-workbook"},
    )
    assert response.status_code == 422
    assert response.json()["diagnostics"][0]["code"] == "CREATION_XLSX_UNREADABLE"
    assert client.get(f"/api/creation-drafts/{creation_draft.id}").json() == before


# ---- 权威预览 -------------------------------------------------------------


def test_preview_returns_group_table_and_sheet_details(
    client: TestClient, creation_draft: DraftFixture, tmp_path: Path
) -> None:
    body = preview_draft(client, creation_draft.id).json()
    assert body["executable"] is True
    assert body["diagnostics"] == []
    assert body["standard_id"] == "szmedi.gas"
    assert body["standard_name"] == "市政燃气施工图"
    assert body["target_path"] == str(tmp_path / "projects" / "新建项目")
    assert body["group_count"] == 1
    assert body["sheet_count"] == 2
    assert body["dwg_count"] == 1
    assert body["numbering"] == {"sequence_field": "subset.sequence", "digits": 2, "start": 1}
    assert body["suffix"] == {"enabled": True, "suffix_type": 1, "unnumbered_keywords": []}
    assert body["sheetset_values"] == {
        "prop-name": "示例工程",
        "prop-major": "燃气",
        "prop-code": "RQ",
    }

    group = body["groups"][0]
    assert group["group_id"] == "group-1"
    assert group["title"] == "平面图"
    assert group["number_range"] == "01-02"
    assert group["sheet_count"] == 2
    assert group["dwg_name"] == "RQ-01-02 平面图.dwg"
    assert group["target_path"] == str(
        tmp_path / "projects" / "新建项目" / "RQ-01-02 平面图.dwg"
    )
    assert group["base_template"] == "templates/a1.dwt"
    assert group["layout_template"] == "templates/a1-layout.dwt"
    assert group["paper_layout"] == "A1"
    assert [sheet["number"] for sheet in group["sheets"]] == ["01", "02"]
    assert [sheet["title"] for sheet in group["sheets"]] == ["平面图 (一)", "平面图 (二)"]
    assert group["sheets"][0]["values"]["prop-label"] == "RQ-01"
    assert group["property_cells"]["prop-label"]["first_value"] == "RQ-01"
    assert group["property_cells"]["prop-label"]["sheets"] == [
        {"number": "01", "value": "RQ-01"},
        {"number": "02", "value": "RQ-02"},
    ]


def test_preview_digest_is_stable_without_changes(
    client: TestClient, creation_draft: DraftFixture
) -> None:
    first = preview_draft(client, creation_draft.id).json()
    second = preview_draft(client, creation_draft.id).json()
    assert first["preview_digest"] == second["preview_digest"]


def test_preview_digest_binds_standard_settings_and_assets(
    client: TestClient, creation_draft: DraftFixture, root: Path
) -> None:
    original = preview_draft(client, creation_draft.id).json()["preview_digest"]

    document_path = published_root(root) / "document.json"
    document_path.write_text(
        document_path.read_text(encoding="utf-8").replace("市政燃气施工图", "市政燃气施工图（改）"),
        encoding="utf-8",
    )
    assert preview_draft(client, creation_draft.id).json()["preview_digest"] != original
    assert preview_draft(client, creation_draft.id).json()["standard_name"] == "市政燃气施工图（改）"

    asset_path = published_root(root) / "templates" / "a1.dwt"
    asset_path.write_bytes(b"base-template-bytes-changed")
    assert preview_draft(client, creation_draft.id).json()["preview_digest"] != original

    service = client.app.state.service
    service.settings.number_suffix_type = 2
    assert preview_draft(client, creation_draft.id).json()["preview_digest"] != original


def test_preview_accepts_existing_empty_target(
    client: TestClient, creation_draft: DraftFixture, tmp_path: Path
) -> None:
    """已存在的空目录是合法目标，但目标状态不同必须得到不同摘要。"""
    missing = preview_draft(client, creation_draft.id).json()
    current = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    target = Path(current["target_path"])
    target.mkdir(parents=True)
    empty = preview_draft(client, creation_draft.id).json()
    assert empty["executable"] is True
    assert empty["preview_digest"] != missing["preview_digest"]


def test_preview_blocks_target_path_occupied_by_file(
    client: TestClient, creation_draft: DraftFixture, tmp_path: Path
) -> None:
    """目标位置存在同名文件时同样不能作为新项目目录。"""
    current = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    target = tmp_path / "occupied" / "新建项目"
    target.parent.mkdir(parents=True)
    target.write_text("x", encoding="utf-8")
    save_creation_draft(
        client,
        creation_draft.id,
        expected_revision=current["revision"],
        target_path=str(target),
        groups=current["groups"],
        sheetset_values=current["sheetset_values"],
    )
    response = preview_draft(client, creation_draft.id)
    assert response.status_code == 422
    assert response.json()["code"] == "CREATION_TARGET_NOT_EMPTY"


def test_preview_digest_binds_target_path(
    client: TestClient, creation_draft: DraftFixture, tmp_path: Path
) -> None:
    original = preview_draft(client, creation_draft.id).json()["preview_digest"]
    current = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    save_creation_draft(
        client,
        creation_draft.id,
        expected_revision=current["revision"],
        target_path=str(tmp_path / "other" / "新建项目"),
        groups=current["groups"],
        sheetset_values=current["sheetset_values"],
    )
    assert preview_draft(client, creation_draft.id).json()["preview_digest"] != original


def test_preview_blocks_empty_group_draft(
    client: TestClient, root: Path, tmp_path: Path
) -> None:
    publish_standard(client, root)
    draft = create_creation_draft(client)
    saved = save_creation_draft(
        client,
        draft["id"],
        expected_revision=draft["revision"],
        target_path=str(tmp_path / "projects" / "新建项目"),
        groups=[],
    )
    assert saved["groups"] == []
    body = preview_draft(client, draft["id"]).json()
    assert body["executable"] is False
    assert body["group_count"] == 0
    assert body["diagnostics"][0]["code"] == "CREATION_GROUPS_EMPTY"
    assert body["diagnostics"][0]["severity"] == "error"


def test_preview_blocks_missing_template_files(
    client: TestClient, root: Path, tmp_path: Path
) -> None:
    """发布后模板文件消失：预览必须显式阻断，不给出可执行结论。"""
    publish_standard(client, root)
    draft = create_creation_draft(client)
    save_creation_draft(
        client,
        draft["id"],
        expected_revision=draft["revision"],
        target_path=str(tmp_path / "projects" / "新建项目"),
    )
    (published_root(root) / "templates" / "a1.dwt").unlink()
    body = preview_draft(client, draft["id"]).json()
    assert body["executable"] is False
    assert [item["code"] for item in body["diagnostics"]] == ["CREATION_ASSET_FILE_MISSING"]


def test_preview_reports_invalid_path_shape_as_diagnostic(
    client: TestClient, creation_draft: DraftFixture
) -> None:
    """路径形状非法由领域诊断定位（不做文件系统判定），不冒泡成 422。"""
    current = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    save_creation_draft(
        client,
        creation_draft.id,
        expected_revision=current["revision"],
        target_path="相对目录\\新建项目",
        groups=current["groups"],
        sheetset_values=current["sheetset_values"],
    )
    body = preview_draft(client, creation_draft.id).json()
    assert body["executable"] is False
    assert [item["code"] for item in body["diagnostics"]] == [
        "CREATION_TARGET_PATH_NOT_ABSOLUTE"
    ]


def test_preview_reports_missing_published_standard(
    client: TestClient, creation_draft: DraftFixture, root: Path
) -> None:
    shutil.rmtree(published_root(root))
    response = preview_draft(client, creation_draft.id)
    assert response.status_code == 404
    assert response.json()["code"] == "CREATION_STANDARD_MISSING"


def test_preview_never_starts_cad(
    client: TestClient, creation_draft: DraftFixture, monkeypatch: pytest.MonkeyPatch
) -> None:
    forbid_cad(monkeypatch)
    body = preview_draft(client, creation_draft.id).json()
    assert body["executable"] is True


def test_draft_save_rejects_derived_and_unknown_input(
    client: TestClient, creation_draft: DraftFixture
) -> None:
    current = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    response = client.put(
        f"/api/creation-drafts/{creation_draft.id}",
        json={
            "expected_revision": current["revision"],
            "step": "review",
            "target_path": current["target_path"],
            "sheetset_values": {**current["sheetset_values"], "prop-code": "伪造派生值"},
            "groups": current["groups"],
        },
    )
    assert response.status_code == 422
    assert response.json()["code"] == "CREATION_DRAFT_INVALID"
    assert client.get(f"/api/creation-drafts/{creation_draft.id}").json() == current


def test_draft_save_rejects_unknown_step(
    client: TestClient, creation_draft: DraftFixture
) -> None:
    current = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    response = client.put(
        f"/api/creation-drafts/{creation_draft.id}",
        json={
            "expected_revision": current["revision"],
            "step": "execute",
            "target_path": current["target_path"],
            "sheetset_values": current["sheetset_values"],
            "groups": current["groups"],
        },
    )
    assert response.status_code == 422


def test_draft_document_shape_is_unchanged_by_preview(
    client: TestClient, creation_draft: DraftFixture
) -> None:
    """预览不得写文件、不得改动草稿（含修订号）。"""
    before = client.get(f"/api/creation-drafts/{creation_draft.id}").json()
    assert preview_draft(client, creation_draft.id).status_code == 200
    assert client.get(f"/api/creation-drafts/{creation_draft.id}").json() == before


def test_standard_document_fixture_is_publishable() -> None:
    """夹具标准必须是可发布文档，避免测试用非法标准掩盖契约问题。"""
    standard = parse_published_standard_document(copy.deepcopy(STANDARD_DOCUMENT))
    assert standard.standard_id == "szmedi.gas"
    assert [asset.asset_id for asset in standard.assets] == ["base-a1", "layout-a1"]
