"""图纸目录 XLSX 生成（PLAN-DB-001 Task 8，SPEC-DB-001 §8）。

固定工作表 ``图纸目录``、第一行 ``图号/图名/DWG/布局``、第二行唯一数据行、
图号为字符串保留前导零；同输入两次生成字节完全一致。
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from dst_builder.infrastructure.catalog.xlsx import (
    SHEET_CATALOG_HEADERS,
    SHEET_CATALOG_SHEET_NAME,
    build_sheet_catalog_xlsx,
    load_sheet_catalog,
)

_ROW = {
    "number": "001",
    "title": "首层平面图",
    "dwg": "001 首层平面图.dwg",
    "layout": "001 首层平面图",
}


def test_fixed_sheet_name_and_headers() -> None:
    data = build_sheet_catalog_xlsx(**_ROW)
    workbook = load_workbook(BytesIO(data))
    assert workbook.sheetnames == [SHEET_CATALOG_SHEET_NAME] == ["图纸目录"]
    worksheet = workbook[SHEET_CATALOG_SHEET_NAME]
    first_row = [worksheet.cell(row=1, column=column).value for column in range(1, 5)]
    assert tuple(first_row) == SHEET_CATALOG_HEADERS == ("图号", "图名", "DWG", "布局")


def test_single_data_row_matches_plan_values() -> None:
    data = build_sheet_catalog_xlsx(**_ROW)
    rows = load_sheet_catalog(data)
    assert rows == [
        ["图号", "图名", "DWG", "布局"],
        ["001", "首层平面图", "001 首层平面图.dwg", "001 首层平面图"],
    ]


def test_number_is_string_cell_preserving_leading_zeros() -> None:
    data = build_sheet_catalog_xlsx(**_ROW)
    workbook = load_workbook(BytesIO(data))
    worksheet = workbook[SHEET_CATALOG_SHEET_NAME]
    cell = worksheet.cell(row=2, column=1)
    assert isinstance(cell.value, str)
    assert cell.value == "001"
    assert cell.data_type == "s"  # 字符串单元格，不是数字


def test_generation_is_byte_deterministic() -> None:
    first = build_sheet_catalog_xlsx(**_ROW)
    second = build_sheet_catalog_xlsx(**_ROW)
    assert first == second
