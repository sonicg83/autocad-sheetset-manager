"""PLAN-DM-020 Task 7：图纸目录候选工作簿生成与回读验证。

覆盖：唯一可见"图纸目录"工作表、首行列名、``freeze_panes == "A2"``、
auto_filter、数据行按传入顺序、全部单元格字符串类型、``001`` 前导零保留、
``=+-@`` 开头不成为公式、空集仅表头、列宽上限与换行、禁止特性（公式/宏/
图表/外部链接/隐藏列/隐藏表/定义名称）以及伪造候选被
``validate_candidate`` 拒绝。
"""

import zipfile
from dataclasses import fields as dataclass_fields
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from dst_manager.extensions.builtin.sheet_catalog.errors import (
    SHEET_CATALOG_BLOCKING,
    SHEET_CATALOG_MESSAGE_KEYS,
    SheetCatalogError,
)
from dst_manager.extensions.builtin.sheet_catalog.workbook import (
    MAX_COLUMN_WIDTH,
    WORKSHEET_NAME,
    WorkbookSummary,
    validate_candidate,
    write_candidate,
)

HEADERS = ("序号", "图号", "图纸名称", "图幅", "备注")
ROWS = (
    ("1", "JG-001", "总平面布置图", "A1", ""),
    ("2", "JG-002", "基础平面图", "A2", "001"),
)


def make_candidate(tmp_path: Path) -> Path:
    path = tmp_path / "candidate.xlsx"
    write_candidate(path, HEADERS, ROWS)
    return path


def reload_candidate(path: Path):
    wb = load_workbook(path, data_only=False, read_only=False)
    return wb, wb[wb.sheetnames[0]]


def save_raw_workbook(path: Path, build) -> None:
    """用 openpyxl 手工构造一个候选文件并落盘（伪造候选用）。"""
    wb = Workbook()
    build(wb)
    wb.save(path)


# ---------------------------------------------------------------- 写入形状


def test_write_candidate_creates_single_visible_sheet_named_catalog(tmp_path):
    path = make_candidate(tmp_path)
    wb, ws = reload_candidate(path)
    assert len(wb.sheetnames) == 1
    assert ws.title == WORKSHEET_NAME
    assert ws.sheet_state == "visible"


def test_write_candidate_pins_header_row_freeze_panes_and_auto_filter(tmp_path):
    path = make_candidate(tmp_path)
    _, ws = reload_candidate(path)
    assert tuple(cell.value for cell in ws[1]) == HEADERS
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_write_candidate_writes_rows_in_order_as_text(tmp_path):
    path = make_candidate(tmp_path)
    _, ws = reload_candidate(path)
    for row_idx, expected in enumerate(ROWS, start=2):
        # 空字符串单元格落盘后读回为空单元格（None）：XLSX 无持久化空串。
        normalized = tuple(value if value != "" else None for value in expected)
        assert tuple(cell.value for cell in ws[row_idx]) == normalized
        for cell in ws[row_idx]:
            if cell.value is None:  # 空值落盘为空单元格，绝不成为公式。
                assert cell.data_type != "f"
            else:
                assert cell.data_type == "s"
    for cell in ws[1]:
        assert cell.data_type == "s"


@pytest.mark.parametrize(
    "raw",
    ["001", "=SUM(A1)", "+1+1", "-2", "@cmd"],
    ids=["leading-zero", "equals", "plus", "minus", "at"],
)
def test_write_candidate_keeps_formula_like_prefixes_as_text(tmp_path, raw):
    path = tmp_path / "candidate.xlsx"
    write_candidate(path, ("图号",), ((raw,),))
    _, ws = reload_candidate(path)
    cell = ws["A2"]
    assert cell.value == raw
    assert cell.data_type == "s"


def test_write_candidate_empty_rows_yields_header_only_workbook(tmp_path):
    path = tmp_path / "empty.xlsx"
    summary = write_candidate(path, HEADERS, ())
    wb, ws = reload_candidate(path)
    assert len(wb.sheetnames) == 1
    assert ws.max_row == 1
    assert tuple(cell.value for cell in ws[1]) == HEADERS
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert summary.data_rows == 0
    assert validate_candidate(path, HEADERS, 0).data_rows == 0


def test_write_candidate_all_empty_rows_roundtrip(tmp_path):
    """全部字段为空的数据行必须保住行数：写出、读回、自检三者一致。"""
    path = tmp_path / "all-empty.xlsx"
    summary = write_candidate(path, HEADERS, (("",) * 5, ("",) * 5, ("",) * 5))
    _wb, ws = reload_candidate(path)
    assert ws.max_row == 4
    assert all(
        cell.value is None for row in ws.iter_rows(min_row=2) for cell in row
    )
    assert summary.data_rows == 3
    assert validate_candidate(path, HEADERS, 3).data_rows == 3


def test_write_candidate_trailing_empty_row_keeps_row_count(tmp_path):
    """行尾全空行读回后行仍存在，data_rows 与输入行数一致。"""
    path = tmp_path / "trailing.xlsx"
    rows = (("1", "JG-001", "总平面布置图", "A1", ""), ("", "", "", "", ""))
    summary = write_candidate(path, HEADERS, rows)
    _, ws = reload_candidate(path)
    assert ws.max_row == 3
    assert tuple(cell.value for cell in ws[3]) == (None,) * 5
    assert summary.data_rows == 2
    assert validate_candidate(path, HEADERS, 2).data_rows == 2


def test_write_candidate_caps_column_width_and_wraps_long_content(tmp_path):
    long_title = "很长的图纸名称" * 12
    path = tmp_path / "wide.xlsx"
    write_candidate(path, HEADERS, (("1", "JG-001", long_title, "A1", ""),))
    _, ws = reload_candidate(path)
    for letter in ("A", "B", "C", "D", "E"):
        width = ws.column_dimensions[letter].width
        assert width is not None
        assert width <= MAX_COLUMN_WIDTH
    assert ws["C2"].alignment.wrap_text is True


def test_write_candidate_contains_no_forbidden_features(tmp_path):
    path = make_candidate(tmp_path)
    wb, ws = reload_candidate(path)
    assert list(wb.defined_names) == []
    assert wb._external_links == []
    assert ws._charts == []
    assert all(not dim.hidden for dim in ws.column_dimensions.values())
    with zipfile.ZipFile(path) as zf:
        names = [name.lower() for name in zf.namelist()]
    for marker in ("vbaproject", "externallink", "chart", "activex"):
        assert not any(marker in name for name in names)


def test_workbook_summary_shape_is_frozen_and_pinned():
    assert [f.name for f in dataclass_fields(WorkbookSummary)] == [
        "worksheet_name",
        "headers",
        "data_rows",
    ]
    summary = WorkbookSummary(worksheet_name="图纸目录", headers=HEADERS, data_rows=2)
    assert summary.worksheet_name == "图纸目录"
    assert summary.headers == HEADERS
    assert summary.data_rows == 2
    with pytest.raises(AttributeError):  # FrozenInstanceError
        summary.data_rows = 3  # type: ignore[misc]


# ---------------------------------------------------------------- 回读校验


def test_validate_candidate_returns_matching_summary(tmp_path):
    path = make_candidate(tmp_path)
    summary = validate_candidate(path, HEADERS, len(ROWS))
    assert isinstance(summary, WorkbookSummary)
    assert summary.worksheet_name == "图纸目录"
    assert summary.headers == HEADERS
    assert summary.data_rows == 2


def test_validate_candidate_rejects_multi_sheet_workbook(tmp_path):
    forged = tmp_path / "multi.xlsx"
    save_raw_workbook(
        forged,
        lambda wb: (
            setattr(wb.active, "title", WORKSHEET_NAME),
            wb.create_sheet("额外表"),
        ),
    )
    with pytest.raises(SheetCatalogError) as excinfo:
        validate_candidate(forged, HEADERS, len(ROWS))
    assert excinfo.value.code == "SHEET_CATALOG_XLSX_INVALID"
    assert excinfo.value.message_key == SHEET_CATALOG_MESSAGE_KEYS[
        "SHEET_CATALOG_XLSX_INVALID"
    ]
    assert SHEET_CATALOG_BLOCKING["SHEET_CATALOG_XLSX_INVALID"] is True


def forge_hidden_sheet(path: Path) -> Path:
    """把唯一工作表改成 hidden（openpyxl 拒绝保存单隐藏表，需在包级伪造）。"""
    with zipfile.ZipFile(path) as zf:
        entries = {info.filename: zf.read(info.filename) for info in zf.infolist()}
    workbook_xml = entries["xl/workbook.xml"].decode("utf-8")
    entries["xl/workbook.xml"] = workbook_xml.replace(
        'state="visible"', 'state="hidden"'
    ).encode("utf-8")
    forged = path.with_name(f"{path.stem}-hidden.xlsx")
    with zipfile.ZipFile(forged, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return forged


def test_validate_candidate_rejects_hidden_sheet(tmp_path):
    path = make_candidate(tmp_path)
    forged = forge_hidden_sheet(path)
    with pytest.raises(SheetCatalogError) as excinfo:
        validate_candidate(forged, HEADERS, len(ROWS))
    assert excinfo.value.code == "SHEET_CATALOG_XLSX_INVALID"


def test_validate_candidate_rejects_wrong_headers(tmp_path):
    path = make_candidate(tmp_path)
    with pytest.raises(SheetCatalogError) as excinfo:
        validate_candidate(path, ("序号", "图号", "名称", "图幅", "备注"), len(ROWS))
    assert excinfo.value.code == "SHEET_CATALOG_XLSX_INVALID"


def test_validate_candidate_rejects_row_count_mismatch(tmp_path):
    path = make_candidate(tmp_path)
    for expected in (len(ROWS) - 1, len(ROWS) + 1):
        with pytest.raises(SheetCatalogError) as excinfo:
            validate_candidate(path, HEADERS, expected)
        assert excinfo.value.code == "SHEET_CATALOG_XLSX_INVALID"


def test_validate_candidate_rejects_formula_cells(tmp_path):
    forged = tmp_path / "formula.xlsx"
    save_raw_workbook(
        forged,
        lambda wb: (
            setattr(wb.active, "title", WORKSHEET_NAME),
            setattr(wb.active, "A1", "图号"),
            setattr(wb.active, "A2", "=SUM(A1)"),
        ),
    )
    with pytest.raises(SheetCatalogError) as excinfo:
        validate_candidate(forged, ("图号",), 1)
    assert excinfo.value.code == "SHEET_CATALOG_XLSX_INVALID"


def test_validate_candidate_rejects_hidden_column(tmp_path):
    forged = tmp_path / "hidden-col.xlsx"
    save_raw_workbook(
        forged,
        lambda wb: (
            setattr(wb.active, "title", WORKSHEET_NAME),
            setattr(wb.active, "A1", "图号"),
            setattr(wb.active.column_dimensions["A"], "hidden", True),
        ),
    )
    with pytest.raises(SheetCatalogError) as excinfo:
        validate_candidate(forged, ("图号",), 0)
    assert excinfo.value.code == "SHEET_CATALOG_XLSX_INVALID"


def inject_zip_entry(path: Path, name: str, payload: bytes) -> Path:
    """把一个额外部件注入已保存的 XLSX 包（伪造宏/外链/图表候选用）。"""
    with zipfile.ZipFile(path) as zf:
        entries = [(info.filename, zf.read(info.filename)) for info in zf.infolist()]
    forged = path.with_name(f"{path.stem}-forged.xlsx")
    with zipfile.ZipFile(forged, "w", zipfile.ZIP_DEFLATED) as zf:
        for entry_name, data in entries:
            zf.writestr(entry_name, data)
        zf.writestr(name, payload)
    return forged


@pytest.mark.parametrize(
    "part_name",
    [
        "xl/vbaProject.bin",
        "xl/externalLinks/externalLink1.xml",
        "xl/charts/chart1.xml",
        "xl/activeX/activeX1.xml",
    ],
)
def test_validate_candidate_rejects_forbidden_package_parts(tmp_path, part_name):
    path = make_candidate(tmp_path)
    forged = inject_zip_entry(path, part_name, b"<forged/>")
    with pytest.raises(SheetCatalogError) as excinfo:
        validate_candidate(forged, HEADERS, len(ROWS))
    assert excinfo.value.code == "SHEET_CATALOG_XLSX_INVALID"


def test_validate_candidate_rejects_unreadable_file(tmp_path):
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_bytes(b"not an xlsx")
    with pytest.raises(SheetCatalogError) as excinfo:
        validate_candidate(corrupt, HEADERS, 0)
    assert excinfo.value.code == "SHEET_CATALOG_XLSX_INVALID"
