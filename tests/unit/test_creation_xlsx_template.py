"""创建模板生成测试：可见表形状、表头去重与隐藏技术元数据（PLAN-DM-036 Task 2）。

覆盖 brief Step 1 的模板形状用例与 Step 4 的隐藏元数据用例：只有 ``SheetSet``
与 ``Sheet`` 两个可见表、固定表头、动态列与固定表头重名时按稳定 ``property_id``
去重、隐藏表记录模板版本/标准身份/列映射/资产映射，可见表不出现内部 ID、派生列
或公式，且枚举与模板选项带 Data Validation。
"""

from io import BytesIO

from creation_xlsx_fixtures import build_creation_template
from openpyxl import load_workbook

from dst_manager.infrastructure.creation_xlsx import (
    LIST_SHEET,
    META_SHEET,
    SHEET_SHEET,
    SHEETSET_SHEET,
)


def test_template_has_legacy_style_two_visible_sheets_and_validation(standard, options) -> None:
    workbook = load_workbook(BytesIO(build_creation_template(standard, options)))
    assert [sheet.title for sheet in workbook if sheet.sheet_state == "visible"] == ["SheetSet", "Sheet"]
    assert workbook["Sheet"].cell(1, 1).value == "图名"
    assert [workbook["Sheet"].cell(1, col).value for col in range(1, 6)] == [
        "图名", "张数", "基础模板", "布局模板", "图幅"
    ]
    assert "派生属性" not in [cell.value for cell in workbook["Sheet"][1]]
    assert len(workbook["Sheet"].data_validations.dataValidation) >= 3


def test_fixed_header_collision_is_disambiguated_by_stable_property_id(standard_with_sheet_property_named_count, options) -> None:
    workbook = load_workbook(BytesIO(build_creation_template(standard_with_sheet_property_named_count, options)))
    headers = [cell.value for cell in workbook["Sheet"][1]]
    assert headers[1] == "张数"
    assert headers.count("张数") == 1
    assert len(headers) == len(set(headers))


def test_template_hidden_metadata_maps_labels_to_stable_ids(standard, options) -> None:
    workbook = load_workbook(BytesIO(build_creation_template(standard, options)))
    assert [sheet.title for sheet in workbook if sheet.sheet_state != "visible"] == [
        META_SHEET,
        LIST_SHEET,
    ]
    visible_values = [
        str(cell.value)
        for sheet in workbook
        if sheet.sheet_state == "visible"
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    # 可见表不出现内部 ID、派生列与任何公式。
    assert not any("prop-" in value or "base-a1" in value or "layout-a1" in value for value in visible_values)
    assert "图签" not in [cell.value for cell in workbook[SHEET_SHEET][1]]
    assert "专业代码" not in [cell.value for cell in workbook[SHEETSET_SHEET]["A"]]
    assert all(
        cell.data_type != "f"
        for sheet in workbook
        for row in sheet.iter_rows()
        for cell in row
    )
    records = {
        (row[0].value, row[1].value): (row[2].value, row[3].value)
        for row in workbook[META_SHEET].iter_rows()
    }
    assert records[("version", "template_version")] == ("1", None)
    assert records[("standard", "standard_id")] == (standard.standard_id, None)
    assert records[("standard", "standard_version")] == (str(standard.version), None)
    assert records[("fixed_column", "A")] == ("title", "图名")
    assert records[("property_column", "G")] == ("prop-part", "分部")
    assert records[("asset", "base-template")] == ("base-a1", "a1.dwt")
    assert records[("asset", "layout-template")] == ("layout-a1", "a1-layout.dwt")
    assert records[("sheetset_target_path", "2")] == (None, "项目保存路径")
    assert records[("sheetset_property", "3")] == ("prop-name", "工程名称")
