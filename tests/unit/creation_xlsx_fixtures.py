"""创建 XLSX 模板与导入测试的共享夹具数据与辅助函数（PLAN-DM-036 Task 2）。

夹具标准与资产候选按模块常量固定，测试辅助函数与测试用例共用同一份定义；
工作簿只在内存 ``BytesIO`` 中往返，不写盘、不依赖本地数据库或上次运行残留。
测试用 openpyxl 直接填写模板，模拟用户填写与「粘贴绕过」两种路径；解析端必须
独立校验，不依赖 Data Validation。
"""

import copy
import zipfile
from io import BytesIO

from openpyxl import load_workbook

from dst_manager.application.creation_import import build_creation_template
from dst_manager.domain.creation import (
    CreationAssetOption,
    CreationDiagnostic,
    CreationImportResult,
)
from dst_manager.domain.standards import (
    DrawingStandard,
    parse_published_standard_document,
)
from dst_manager.infrastructure.creation_xlsx import (
    META_SHEET,
    SHEET_SHEET,
    SHEETSET_SHEET,
)

# 最小可发布标准：两个普通 sheetset 属性（文本带默认值、枚举带默认值）、一个派生
# 映射属性、一个普通 sheet 文本属性、一个普通 sheet 枚举属性与一个派生组合属性。
STANDARD_DOCUMENT: dict[str, object] = {
    "schema_version": 1,
    "standard_id": "szmedi.gas",
    "version": "2.1.0",
    "name": "市政燃气施工图",
    "supported_cad_versions": ["2016", "2020"],
    "properties": [
        {
            "property_id": "prop-name",
            "name": "工程名称",
            "scope": "sheetset",
            "kind": "text",
            "required": True,
            "default_value": "默认工程",
        },
        {
            "property_id": "prop-major",
            "name": "专业",
            "scope": "sheetset",
            "kind": "enum",
            "required": True,
            "default_value": "燃气",
            "enum_items": [
                {"item_id": "enum-gas", "value": "燃气"},
                {"item_id": "enum-oil", "value": "燃油"},
            ],
        },
        {
            "property_id": "prop-code",
            "name": "专业代码",
            "scope": "sheetset",
            "kind": "mapping",
            "source_property_id": "prop-major",
            "mapping": [
                {"item_id": "enum-gas", "value": "RQ"},
                {"item_id": "enum-oil", "value": "RY"},
            ],
            "confirmed_source_items": [["enum-gas", "燃气"], ["enum-oil", "燃油"]],
        },
        {
            "property_id": "prop-stage",
            "name": "设计阶段",
            "scope": "sheet",
            "kind": "text",
            "default_value": "施工图",
        },
        {
            "property_id": "prop-part",
            "name": "分部",
            "scope": "sheet",
            "kind": "enum",
            "enum_items": [
                {"item_id": "enum-part-a", "value": "A 段"},
                {"item_id": "enum-part-b", "value": "B 段"},
            ],
        },
        {
            "property_id": "prop-label",
            "name": "图签",
            "scope": "sheet",
            "kind": "composition",
            "segments": [
                {"property_id": "prop-code"},
                {"literal": "-"},
                {"system_field": "sheet.number"},
            ],
        },
    ],
    "dwg_naming": {
        "segments": [
            {"property_id": "prop-code"},
            {"literal": "-"},
            {"system_field": "subset.scope"},
            {"literal": " "},
            {"system_field": "subset.name"},
        ]
    },
    "assets": [
        {
            "asset_id": "base-a1",
            "kind": "base-template",
            "files": [{"path": "templates/a1.dwt", "role": ""}],
        },
        {
            "asset_id": "layout-a1",
            "kind": "layout-template",
            "files": [{"path": "templates/a1-layout.dwt", "role": "A1"}],
        },
    ],
    "numbering": {"sequence_field": "subset.sequence", "digits": 2},
}

#: 夹具标准与资产候选按模块常量固定：测试辅助函数与测试用例共用同一份定义。
STANDARD = parse_published_standard_document(STANDARD_DOCUMENT)

_COUNT_PROPERTY_DOCUMENT = copy.deepcopy(STANDARD_DOCUMENT)
_COUNT_PROPERTY_PROPERTIES = _COUNT_PROPERTY_DOCUMENT["properties"]
assert isinstance(_COUNT_PROPERTY_PROPERTIES, list)
_COUNT_PROPERTY_PROPERTIES.append(
    {
        "property_id": "prop-count",
        "name": "张数",
        "scope": "sheet",
        "kind": "text",
        "default_value": "",
    }
)
STANDARD_WITH_COUNT_PROPERTY = parse_published_standard_document(_COUNT_PROPERTY_DOCUMENT)

#: 资产候选标签按受控包内文件名生成（同类内唯一由构造方保证）。
ASSET_OPTIONS: tuple[CreationAssetOption, ...] = (
    CreationAssetOption(asset_id="base-a1", kind="base-template", label="a1.dwt"),
    CreationAssetOption(
        asset_id="layout-a1", kind="layout-template", label="a1-layout.dwt", layouts=("A1",)
    ),
)

DEFAULT_SHEETSET_PATH = r"C:\Projects\新建项目"


def group_row(**overrides: object) -> dict[str, object]:
    """一行合法图纸组输入（按可见表头键控）；``overrides`` 用于构造非法输入。"""
    row: dict[str, object] = {
        "图名": "平面图",
        "张数": 1,
        "基础模板": "a1.dwt",
        "布局模板": "a1-layout.dwt",
        "图幅": "A1",
        "设计阶段": "施工图",
        "分部": "A 段",
    }
    row.update(overrides)
    return row


def _label_row(worksheet, label: str) -> int:
    for row in range(1, worksheet.max_row + 1):
        if worksheet.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"模板缺少行标签 {label!r}")


def _header_column(worksheet, header: str) -> int:
    for column in range(1, worksheet.max_column + 1):
        if worksheet.cell(row=1, column=column).value == header:
            return column
    raise AssertionError(f"模板缺少表头 {header!r}")


def fill_template(
    standard: DrawingStandard,
    options: tuple[CreationAssetOption, ...],
    *,
    sheetset: dict[str, object] | None = None,
    rows: tuple[dict[str, object], ...] = (),
    mutate=None,
) -> bytes:
    """在模板上按可见属性名/表头填值，返回工作簿字节。

    ``mutate`` 用于模拟粘贴绕过与隐藏表篡改：直接改工作簿对象后再保存。
    """
    workbook = load_workbook(BytesIO(build_creation_template(standard, options)))
    sheetset_sheet = workbook[SHEETSET_SHEET]
    sheet_sheet = workbook[SHEET_SHEET]
    for label, value in (sheetset or {}).items():
        sheetset_sheet.cell(row=_label_row(sheetset_sheet, label), column=2).value = value
    for offset, row in enumerate(rows):
        for header, value in row.items():
            sheet_sheet.cell(row=2 + offset, column=_header_column(sheet_sheet, header)).value = value
    if mutate is not None:
        mutate(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def meta_cell(workbook, record: str, key: str):
    """隐藏技术表里 ``(record, key)`` 记录的值单元格（C 列）。"""
    meta = workbook[META_SHEET]
    for row in range(1, meta.max_row + 1):
        if meta.cell(row=row, column=1).value == record and str(
            meta.cell(row=row, column=2).value
        ) == key:
            return meta.cell(row=row, column=3)
    raise AssertionError(f"隐藏技术表缺少记录 {record!r}/{key!r}")


def meta_label_cell(workbook, record: str, label: str):
    """隐藏技术表里按 ``label`` 定位记录的值单元格（C 列）。"""
    meta = workbook[META_SHEET]
    for row in range(1, meta.max_row + 1):
        if meta.cell(row=row, column=1).value == record and meta.cell(
            row=row, column=4
        ).value == label:
            return meta.cell(row=row, column=3)
    raise AssertionError(f"隐藏技术表缺少标签 {record!r}/{label!r}")


def workbook_with_invalid_enum() -> bytes:
    """粘贴绕过 Data Validation：Sheet 行内写入不属于枚举列表的分部值。"""
    return fill_template(
        STANDARD,
        ASSET_OPTIONS,
        sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
        rows=(group_row(分部="C 段"),),
    )


def valid_group_workbook(*, path: str, count: int, title: str = "平面图") -> bytes:
    """一份合法工作簿：给定最终路径、图名与张数，其余取合法值。"""
    return fill_template(
        STANDARD,
        ASSET_OPTIONS,
        sheetset={"项目保存路径": path},
        rows=(group_row(图名=title, 张数=count),),
    )


def diagnostic(result: CreationImportResult, code: str) -> CreationDiagnostic:
    """取指定错误码的诊断；缺失时列出实际诊断便于定位。"""
    matches = [item for item in result.diagnostics if item.code == code]
    assert matches, f"缺少诊断 {code}：{[item.code for item in result.diagnostics]}"
    return matches[0]


def with_extra_package_part(data: bytes, name: str) -> bytes:
    """在 XLSX 包内追加一个部件，用于模拟外部链接/宏。"""
    source = zipfile.ZipFile(BytesIO(data))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr(name, b"<extra/>")
    return buffer.getvalue()
