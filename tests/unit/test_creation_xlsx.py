"""双工作表 XLSX 模板与全量结构导入测试（PLAN-DM-036 Task 2）。

覆盖：模板形状（仅 ``SheetSet``/``Sheet`` 两个可见表、固定表头、隐藏技术元数据、
Data Validation）、合法工作簿按行序往返，以及粘贴绕过 Data Validation 的非法枚举、
公式单元格、外部链接与宏、未知/重复/缺失表头、篡改的隐藏映射、失效资产、图幅与
布局模板不匹配、重复图名、非法张数、危险路径与受控规模超限的整批拒绝。

测试用 openpyxl 直接填写模板，模拟用户填写与「粘贴绕过」两种路径；解析端必须
独立校验，不依赖 Data Validation。
"""

import copy
import zipfile
from io import BytesIO

import pytest
from openpyxl import load_workbook

from dst_manager.application.creation_import import (
    build_creation_template,
    parse_creation_workbook,
)
from dst_manager.domain.creation import (
    MAX_TARGET_PATH_LENGTH,
    CreationAssetOption,
    CreationDiagnostic,
    CreationImportResult,
    validate_creation_target_path,
)
from dst_manager.domain.standards import (
    DrawingStandard,
    parse_published_standard_document,
)
from dst_manager.infrastructure.creation_xlsx import (
    LIST_SHEET,
    MAX_SHEET_COLUMNS,
    MAX_SHEET_ROWS,
    META_SHEET,
    SHEET_SHEET,
    SHEETSET_SHEET,
)

# 最小可发布标准：两个普通 sheetset 属性（文本带默认值、枚举带默认值）、一个派生
# 映射属性、一个普通 sheet 文本属性、一个普通 sheet 枚举属性与一个派生组合属性。
STANDARD_DOCUMENT: dict[str, object] = {
    "schema_version": 1,
    "standard_id": "szmedi.gas",
    "version": "2.1.0",
    "name": "市政燃气施工图",
    "supported_cad_versions": ["2016", "2020"],
    "properties": [
        {
            "property_id": "prop-name",
            "name": "工程名称",
            "scope": "sheetset",
            "kind": "text",
            "required": True,
            "default_value": "默认工程",
        },
        {
            "property_id": "prop-major",
            "name": "专业",
            "scope": "sheetset",
            "kind": "enum",
            "required": True,
            "default_value": "燃气",
            "enum_items": [
                {"item_id": "enum-gas", "value": "燃气"},
                {"item_id": "enum-oil", "value": "燃油"},
            ],
        },
        {
            "property_id": "prop-code",
            "name": "专业代码",
            "scope": "sheetset",
            "kind": "mapping",
            "source_property_id": "prop-major",
            "mapping": [
                {"item_id": "enum-gas", "value": "RQ"},
                {"item_id": "enum-oil", "value": "RY"},
            ],
            "confirmed_source_items": [["enum-gas", "燃气"], ["enum-oil", "燃油"]],
        },
        {
            "property_id": "prop-stage",
            "name": "设计阶段",
            "scope": "sheet",
            "kind": "text",
            "default_value": "施工图",
        },
        {
            "property_id": "prop-part",
            "name": "分部",
            "scope": "sheet",
            "kind": "enum",
            "enum_items": [
                {"item_id": "enum-part-a", "value": "A 段"},
                {"item_id": "enum-part-b", "value": "B 段"},
            ],
        },
        {
            "property_id": "prop-label",
            "name": "图签",
            "scope": "sheet",
            "kind": "composition",
            "segments": [
                {"property_id": "prop-code"},
                {"literal": "-"},
                {"system_field": "sheet.number"},
            ],
        },
    ],
    "dwg_naming": {
        "segments": [
            {"property_id": "prop-code"},
            {"literal": "-"},
            {"system_field": "subset.scope"},
            {"literal": " "},
            {"system_field": "subset.name"},
        ]
    },
    "assets": [
        {
            "asset_id": "base-a1",
            "kind": "base-template",
            "files": [{"path": "templates/a1.dwt", "role": ""}],
        },
        {
            "asset_id": "layout-a1",
            "kind": "layout-template",
            "files": [{"path": "templates/a1-layout.dwt", "role": "A1"}],
        },
    ],
    "numbering": {"sequence_field": "subset.sequence", "digits": 2},
}

#: 夹具标准与资产候选按模块常量固定：测试辅助函数与夹具共用同一份定义。
STANDARD = parse_published_standard_document(STANDARD_DOCUMENT)

_COUNT_PROPERTY_DOCUMENT = copy.deepcopy(STANDARD_DOCUMENT)
_COUNT_PROPERTY_PROPERTIES = _COUNT_PROPERTY_DOCUMENT["properties"]
assert isinstance(_COUNT_PROPERTY_PROPERTIES, list)
_COUNT_PROPERTY_PROPERTIES.append(
    {
        "property_id": "prop-count",
        "name": "张数",
        "scope": "sheet",
        "kind": "text",
        "default_value": "",
    }
)
STANDARD_WITH_COUNT_PROPERTY = parse_published_standard_document(_COUNT_PROPERTY_DOCUMENT)

#: 资产候选标签按受控包内文件名生成（同类内唯一由构造方保证）。
ASSET_OPTIONS: tuple[CreationAssetOption, ...] = (
    CreationAssetOption(asset_id="base-a1", kind="base-template", label="a1.dwt"),
    CreationAssetOption(
        asset_id="layout-a1", kind="layout-template", label="a1-layout.dwt", layouts=("A1",)
    ),
)

DEFAULT_SHEETSET_PATH = r"C:\Projects\新建项目"


@pytest.fixture
def standard() -> DrawingStandard:
    return STANDARD


@pytest.fixture
def options() -> tuple[CreationAssetOption, ...]:
    return ASSET_OPTIONS


@pytest.fixture
def standard_with_sheet_property_named_count() -> DrawingStandard:
    return STANDARD_WITH_COUNT_PROPERTY


# ---- 测试辅助 ------------------------------------------------------------


def group_row(**overrides: object) -> dict[str, object]:
    """一行合法图纸组输入（按可见表头键控）；``overrides`` 用于构造非法输入。"""
    row: dict[str, object] = {
        "图名": "平面图",
        "张数": 1,
        "基础模板": "a1.dwt",
        "布局模板": "a1-layout.dwt",
        "图幅": "A1",
        "设计阶段": "施工图",
        "分部": "A 段",
    }
    row.update(overrides)
    return row


def _label_row(worksheet, label: str) -> int:
    for row in range(1, worksheet.max_row + 1):
        if worksheet.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"模板缺少行标签 {label!r}")


def _header_column(worksheet, header: str) -> int:
    for column in range(1, worksheet.max_column + 1):
        if worksheet.cell(row=1, column=column).value == header:
            return column
    raise AssertionError(f"模板缺少表头 {header!r}")


def fill_template(
    standard: DrawingStandard,
    options: tuple[CreationAssetOption, ...],
    *,
    sheetset: dict[str, object] | None = None,
    rows: tuple[dict[str, object], ...] = (),
    mutate=None,
) -> bytes:
    """在模板上按可见属性名/表头填值，返回工作簿字节。

    ``mutate`` 用于模拟粘贴绕过与隐藏表篡改：直接改工作簿对象后再保存。
    """
    workbook = load_workbook(BytesIO(build_creation_template(standard, options)))
    sheetset_sheet = workbook[SHEETSET_SHEET]
    sheet_sheet = workbook[SHEET_SHEET]
    for label, value in (sheetset or {}).items():
        sheetset_sheet.cell(row=_label_row(sheetset_sheet, label), column=2).value = value
    for offset, row in enumerate(rows):
        for header, value in row.items():
            sheet_sheet.cell(row=2 + offset, column=_header_column(sheet_sheet, header)).value = value
    if mutate is not None:
        mutate(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def meta_cell(workbook, record: str, key: str):
    """隐藏技术表里 ``(record, key)`` 记录的值单元格（C 列）。"""
    meta = workbook[META_SHEET]
    for row in range(1, meta.max_row + 1):
        if meta.cell(row=row, column=1).value == record and str(
            meta.cell(row=row, column=2).value
        ) == key:
            return meta.cell(row=row, column=3)
    raise AssertionError(f"隐藏技术表缺少记录 {record!r}/{key!r}")


def meta_label_cell(workbook, record: str, label: str):
    """隐藏技术表里按 ``label`` 定位记录的值单元格（C 列）。"""
    meta = workbook[META_SHEET]
    for row in range(1, meta.max_row + 1):
        if meta.cell(row=row, column=1).value == record and meta.cell(
            row=row, column=4
        ).value == label:
            return meta.cell(row=row, column=3)
    raise AssertionError(f"隐藏技术表缺少标签 {record!r}/{label!r}")


def workbook_with_invalid_enum() -> bytes:
    """粘贴绕过 Data Validation：Sheet 行内写入不属于枚举列表的分部值。"""
    return fill_template(
        STANDARD,
        ASSET_OPTIONS,
        sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
        rows=(group_row(分部="C 段"),),
    )


def valid_group_workbook(*, path: str, count: int, title: str = "平面图") -> bytes:
    """一份合法工作簿：给定最终路径、图名与张数，其余取合法值。"""
    return fill_template(
        STANDARD,
        ASSET_OPTIONS,
        sheetset={"项目保存路径": path},
        rows=(group_row(图名=title, 张数=count),),
    )


def diagnostic(result: CreationImportResult, code: str) -> CreationDiagnostic:
    """取指定错误码的诊断；缺失时列出实际诊断便于定位。"""
    matches = [item for item in result.diagnostics if item.code == code]
    assert matches, f"缺少诊断 {code}：{[item.code for item in result.diagnostics]}"
    return matches[0]


def with_extra_package_part(data: bytes, name: str) -> bytes:
    """在 XLSX 包内追加一个部件，用于模拟外部链接/宏。"""
    source = zipfile.ZipFile(BytesIO(data))
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr(name, b"<extra/>")
    return buffer.getvalue()


# ---- Step 1：模板形状与绕过验证 ------------------------------------------


def test_template_has_legacy_style_two_visible_sheets_and_validation(standard, options) -> None:
    workbook = load_workbook(BytesIO(build_creation_template(standard, options)))
    assert [sheet.title for sheet in workbook if sheet.sheet_state == "visible"] == ["SheetSet", "Sheet"]
    assert workbook["Sheet"].cell(1, 1).value == "图名"
    assert [workbook["Sheet"].cell(1, col).value for col in range(1, 6)] == [
        "图名", "张数", "基础模板", "布局模板", "图幅"
    ]
    assert "派生属性" not in [cell.value for cell in workbook["Sheet"][1]]
    assert len(workbook["Sheet"].data_validations.dataValidation) >= 3


def test_fixed_header_collision_is_disambiguated_by_stable_property_id(standard_with_sheet_property_named_count, options) -> None:
    workbook = load_workbook(BytesIO(build_creation_template(standard_with_sheet_property_named_count, options)))
    headers = [cell.value for cell in workbook["Sheet"][1]]
    assert headers[1] == "张数"
    assert headers.count("张数") == 1
    assert len(headers) == len(set(headers))


def test_pasted_invalid_enum_rejects_entire_workbook(standard, options) -> None:
    result = parse_creation_workbook(workbook_with_invalid_enum(), standard, options)
    assert result.value is None
    assert (result.diagnostics[0].sheet, result.diagnostics[0].row) == ("Sheet", 2)


def test_xlsx_final_path_and_group_count_round_trip(standard, options) -> None:
    result = parse_creation_workbook(valid_group_workbook(path=r"C:\Projects\道路工程", count=3), standard, options)
    assert result.value.target_path == r"C:\Projects\道路工程"
    assert len(result.value.groups) == 1
    assert result.value.groups[0].count == 3


# ---- Step 4：模板技术元数据与合法往返 ------------------------------------


def test_template_hidden_metadata_maps_labels_to_stable_ids(standard, options) -> None:
    workbook = load_workbook(BytesIO(build_creation_template(standard, options)))
    assert [sheet.title for sheet in workbook if sheet.sheet_state != "visible"] == [
        META_SHEET,
        LIST_SHEET,
    ]
    visible_values = [
        str(cell.value)
        for sheet in workbook
        if sheet.sheet_state == "visible"
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    # 可见表不出现内部 ID、派生列与任何公式。
    assert not any("prop-" in value or "base-a1" in value or "layout-a1" in value for value in visible_values)
    assert "图签" not in [cell.value for cell in workbook[SHEET_SHEET][1]]
    assert "专业代码" not in [cell.value for cell in workbook[SHEETSET_SHEET]["A"]]
    assert all(
        cell.data_type != "f"
        for sheet in workbook
        for row in sheet.iter_rows()
        for cell in row
    )
    records = {
        (row[0].value, row[1].value): (row[2].value, row[3].value)
        for row in workbook[META_SHEET].iter_rows()
    }
    assert records[("version", "template_version")] == ("1", None)
    assert records[("standard", "standard_id")] == (standard.standard_id, None)
    assert records[("standard", "standard_version")] == (standard.version, None)
    assert records[("fixed_column", "A")] == ("title", "图名")
    assert records[("property_column", "G")] == ("prop-part", "分部")
    assert records[("asset", "base-template")] == ("base-a1", "a1.dwt")
    assert records[("asset", "layout-template")] == ("layout-a1", "a1-layout.dwt")
    assert records[("sheetset_target_path", "2")] == (None, "项目保存路径")
    assert records[("sheetset_property", "3")] == ("prop-name", "工程名称")


def test_valid_workbook_round_trips_groups_in_row_order(standard, options) -> None:
    data = fill_template(
        standard,
        options,
        sheetset={"项目保存路径": r"C:\Projects\道路工程"},
        rows=(
            group_row(图名="平面图", 张数=3),
            {},
            group_row(图名="剖面图", 设计阶段=""),
        ),
    )
    result = parse_creation_workbook(data, standard, options)
    assert result.diagnostics == ()
    value = result.value
    assert value.target_path == r"C:\Projects\道路工程"
    assert value.sheetset_values == {"prop-name": "默认工程", "prop-major": "燃气"}
    assert [group.title for group in value.groups] == ["平面图", "剖面图"]
    assert [group.count for group in value.groups] == [3, 1]
    assert [group.created_order for group in value.groups] == [1, 2]
    assert len({group.group_id for group in value.groups}) == 2
    # 一行一组：张数不展开为逐张输入，普通 sheet 属性每个键都存在。
    assert value.groups[0].sheet_values == {"prop-stage": "施工图", "prop-part": "A 段"}
    assert value.groups[1].sheet_values == {"prop-stage": "", "prop-part": "A 段"}


def test_imported_value_matches_ui_input_shape(standard, options) -> None:
    result = parse_creation_workbook(
        valid_group_workbook(path=r"C:\Projects\道路工程", count=3), standard, options
    )
    value = result.value
    group = value.groups[0]
    assert value.target_path == r"C:\Projects\道路工程"
    assert value.sheetset_values == {"prop-name": "默认工程", "prop-major": "燃气"}
    assert (group.title, group.count) == ("平面图", 3)
    assert (group.base_asset_id, group.layout_asset_id, group.paper_layout) == (
        "base-a1",
        "layout-a1",
        "A1",
    )
    assert group.sheet_values == {"prop-stage": "施工图", "prop-part": "A 段"}


# ---- Step 4：公式、外部链接与不可读输入 ----------------------------------


def test_formula_cell_is_rejected_with_cell_location(standard, options) -> None:
    data = fill_template(
        standard,
        options,
        sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
        rows=(group_row(),),
        mutate=lambda workbook: setattr(
            workbook[SHEET_SHEET].cell(row=2, column=2), "value", "=1+1"
        ),
    )
    result = parse_creation_workbook(data, standard, options)
    assert result.value is None
    found = diagnostic(result, "CREATION_XLSX_FORMULA_FORBIDDEN")
    assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 2, "B")


def test_external_link_and_macro_parts_are_rejected(standard, options) -> None:
    valid = valid_group_workbook(path=DEFAULT_SHEETSET_PATH, count=1)
    linked = parse_creation_workbook(
        with_extra_package_part(valid, "xl/externalLinks/externalLink1.xml"), standard, options
    )
    assert linked.value is None
    assert diagnostic(linked, "CREATION_XLSX_EXTERNAL_LINK_FORBIDDEN").code
    macro = parse_creation_workbook(
        with_extra_package_part(valid, "xl/vbaProject.bin"), standard, options
    )
    assert macro.value is None
    assert diagnostic(macro, "CREATION_XLSX_MACRO_FORBIDDEN").code


def test_unreadable_payload_is_rejected(standard, options) -> None:
    result = parse_creation_workbook("不是工作簿".encode(), standard, options)
    assert result.value is None
    assert diagnostic(result, "CREATION_XLSX_UNREADABLE").code


# ---- Step 4：隐藏元数据篡改与标准身份 ------------------------------------


def test_tampered_standard_identity_in_metadata_is_rejected(standard, options) -> None:
    data = fill_template(
        standard,
        options,
        sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
        rows=(group_row(),),
        mutate=lambda workbook: setattr(
            meta_cell(workbook, "standard", "standard_version"), "value", "9.9.9"
        ),
    )
    result = parse_creation_workbook(data, standard, options)
    assert result.value is None
    assert diagnostic(result, "CREATION_XLSX_STANDARD_MISMATCH").sheet == META_SHEET


def test_tampered_property_column_mapping_is_rejected(standard, options) -> None:
    def tamper(property_id: str):
        return lambda workbook: setattr(
            meta_label_cell(workbook, "property_column", "分部"), "value", property_id
        )

    for property_id in ("prop-label", "prop-unknown"):
        data = fill_template(
            standard,
            options,
            sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
            rows=(group_row(),),
            mutate=tamper(property_id),
        )
        result = parse_creation_workbook(data, standard, options)
        assert result.value is None
        found = diagnostic(result, "CREATION_XLSX_METADATA_INVALID")
        assert found.sheet == META_SHEET


# ---- Step 4：资产、图幅与行输入校验 --------------------------------------


def test_stale_asset_option_is_rejected(standard, options) -> None:
    data = valid_group_workbook(path=DEFAULT_SHEETSET_PATH, count=1)
    result = parse_creation_workbook(data, standard, options[1:])
    assert result.value is None
    found = diagnostic(result, "CREATION_XLSX_ASSET_INVALID")
    assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 2, "C")


def test_layout_template_and_paper_layout_mismatch_is_rejected(standard, options) -> None:
    wrong_layout = fill_template(
        standard,
        options,
        sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
        rows=(group_row(图幅="A0"),),
    )
    result = parse_creation_workbook(wrong_layout, standard, options)
    assert result.value is None
    found = diagnostic(result, "CREATION_XLSX_PAPER_LAYOUT_INVALID")
    assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 2, "E")

    swapped = fill_template(
        standard,
        options,
        sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
        rows=(group_row(基础模板="a1-layout.dwt"),),
    )
    result = parse_creation_workbook(swapped, standard, options)
    assert result.value is None
    found = diagnostic(result, "CREATION_XLSX_ASSET_INVALID")
    assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 2, "C")


def test_duplicate_group_title_is_rejected(standard, options) -> None:
    data = fill_template(
        standard,
        options,
        sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
        rows=(group_row(图名="平面图"), group_row(图名=" 平面图 ")),
    )
    result = parse_creation_workbook(data, standard, options)
    assert result.value is None
    found = diagnostic(result, "CREATION_XLSX_TITLE_DUPLICATE")
    assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 3, "A")


def test_invalid_group_count_is_rejected(standard, options) -> None:
    for count in ("0", "-1", "3.5", "三", ""):
        data = fill_template(
            standard,
            options,
            sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
            rows=(group_row(张数=count),),
        )
        result = parse_creation_workbook(data, standard, options)
        assert result.value is None, count
        found = diagnostic(result, "CREATION_XLSX_COUNT_INVALID")
        assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 2, "B")


def test_unsafe_target_path_is_rejected(standard, options) -> None:
    cases = {
        r"Projects\道路工程": "CREATION_TARGET_PATH_NOT_ABSOLUTE",
        r"C:道路工程": "CREATION_TARGET_PATH_DRIVE_INVALID",
        r"C:\Projects\道路<工程": "CREATION_TARGET_PATH_CHARACTER_INVALID",
        r"C:\Projects\con": "CREATION_TARGET_PATH_RESERVED_DEVICE",
        r"C:\Projects\道路工程.": "CREATION_TARGET_PATH_TRAILING_CHARACTER",
        r"C:\Projects\..\其他": "CREATION_TARGET_PATH_SEGMENT_INVALID",
        "C:\\Projects\\道路工程\\": "CREATION_TARGET_PATH_SEGMENT_INVALID",
        "": "CREATION_TARGET_PATH_EMPTY",
    }
    for path, code in cases.items():
        data = fill_template(
            standard,
            options,
            sheetset={"项目保存路径": path},
            rows=(group_row(),),
        )
        result = parse_creation_workbook(data, standard, options)
        assert result.value is None, path
        found = diagnostic(result, code)
        assert (found.sheet, found.row, found.column) == (SHEETSET_SHEET, 2, "B")


# ---- Step 4：表头与受控规模 ----------------------------------------------


def test_header_duplicate_missing_and_unknown_are_rejected(standard, options) -> None:
    def data_with(mutate) -> bytes:
        return fill_template(
            standard,
            options,
            sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
            rows=(group_row(),),
            mutate=mutate,
        )

    duplicated = parse_creation_workbook(
        data_with(
            lambda workbook: setattr(workbook[SHEET_SHEET].cell(row=1, column=2), "value", "图名")
        ),
        standard,
        options,
    )
    assert duplicated.value is None
    assert (diagnostic(duplicated, "CREATION_XLSX_HEADER_DUPLICATE").column) == "B"

    missing = parse_creation_workbook(
        data_with(
            lambda workbook: setattr(workbook[SHEET_SHEET].cell(row=1, column=2), "value", None)
        ),
        standard,
        options,
    )
    assert missing.value is None
    found = diagnostic(missing, "CREATION_XLSX_HEADER_MISSING")
    assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 1, "B")

    unknown = parse_creation_workbook(
        data_with(
            lambda workbook: setattr(workbook[SHEET_SHEET].cell(row=1, column=8), "value", "多余列")
        ),
        standard,
        options,
    )
    assert unknown.value is None
    found = diagnostic(unknown, "CREATION_XLSX_HEADER_UNKNOWN")
    assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 1, "H")


def test_unexpected_content_outside_plan_is_rejected(standard, options) -> None:
    sheet_extra = parse_creation_workbook(
        fill_template(
            standard,
            options,
            sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
            rows=(group_row(),),
            mutate=lambda workbook: setattr(
                workbook[SHEET_SHEET].cell(row=2, column=8), "value", "越界内容"
            ),
        ),
        standard,
        options,
    )
    assert sheet_extra.value is None
    found = diagnostic(sheet_extra, "CREATION_XLSX_CELL_UNEXPECTED")
    assert (found.sheet, found.row, found.column) == (SHEET_SHEET, 2, "H")

    sheetset_extra = parse_creation_workbook(
        fill_template(
            standard,
            options,
            sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
            rows=(group_row(),),
            mutate=lambda workbook: setattr(
                workbook[SHEETSET_SHEET].cell(row=7, column=1), "value", "多余行"
            ),
        ),
        standard,
        options,
    )
    assert sheetset_extra.value is None
    found = diagnostic(sheetset_extra, "CREATION_XLSX_CELL_UNEXPECTED")
    assert (found.sheet, found.row, found.column) == (SHEETSET_SHEET, 7, "A")


def test_visible_sheet_set_is_enforced_on_import(standard, options) -> None:
    def data_with(mutate) -> bytes:
        return fill_template(
            standard,
            options,
            sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
            rows=(group_row(),),
            mutate=mutate,
        )

    extra = parse_creation_workbook(
        data_with(lambda workbook: workbook.create_sheet("Sheet3")), standard, options
    )
    assert extra.value is None
    assert diagnostic(extra, "CREATION_XLSX_SHEET_INVALID").sheet == "Sheet3"

    removed = parse_creation_workbook(
        data_with(lambda workbook: workbook.remove(workbook[SHEET_SHEET])), standard, options
    )
    assert removed.value is None
    assert diagnostic(removed, "CREATION_XLSX_SHEET_INVALID").sheet == SHEET_SHEET

    metadata_visible = parse_creation_workbook(
        data_with(lambda workbook: setattr(workbook[META_SHEET], "sheet_state", "visible")),
        standard,
        options,
    )
    assert metadata_visible.value is None
    assert diagnostic(metadata_visible, "CREATION_XLSX_METADATA_INVALID").sheet == META_SHEET


def test_oversized_workbook_is_rejected(standard, options) -> None:
    rows = parse_creation_workbook(
        fill_template(
            standard,
            options,
            sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
            rows=(group_row(),),
            mutate=lambda workbook: setattr(
                workbook[SHEET_SHEET].cell(row=MAX_SHEET_ROWS + 1, column=1), "value", "超限"
            ),
        ),
        standard,
        options,
    )
    assert rows.value is None
    found = diagnostic(rows, "CREATION_XLSX_SCALE_EXCEEDED")
    assert (found.sheet, found.row) == (SHEET_SHEET, MAX_SHEET_ROWS + 1)

    columns = parse_creation_workbook(
        fill_template(
            standard,
            options,
            sheetset={"项目保存路径": DEFAULT_SHEETSET_PATH},
            rows=(group_row(),),
            mutate=lambda workbook: setattr(
                workbook[SHEET_SHEET].cell(row=1, column=MAX_SHEET_COLUMNS + 1), "value", "越界"
            ),
        ),
        standard,
        options,
    )
    assert columns.value is None
    found = diagnostic(columns, "CREATION_XLSX_SCALE_EXCEEDED")
    assert found.sheet == SHEET_SHEET
    assert found.column is not None


# ---- 目标路径领域校验（供 Task 3/4/6 复用） ------------------------------


def test_validate_creation_target_path_accepts_absolute_windows_paths() -> None:
    assert validate_creation_target_path(r"C:\Projects\道路工程") == ()
    assert validate_creation_target_path("D:/Projects/项目") == ()
    assert validate_creation_target_path(r"C:\Projects\新建项目 2") == ()


def test_validate_creation_target_path_rejects_unsafe_values() -> None:
    cases = {
        "": "CREATION_TARGET_PATH_EMPTY",
        "   ": "CREATION_TARGET_PATH_EMPTY",
        r"Projects\道路工程": "CREATION_TARGET_PATH_NOT_ABSOLUTE",
        r"\\server\share\项目": "CREATION_TARGET_PATH_NOT_ABSOLUTE",
        r"C:道路工程": "CREATION_TARGET_PATH_DRIVE_INVALID",
        r"C:\Projects\道路|工程": "CREATION_TARGET_PATH_CHARACTER_INVALID",
        r"C:\Projects\道路工程 ": "CREATION_TARGET_PATH_TRAILING_CHARACTER",
        r"C:\Projects\nul": "CREATION_TARGET_PATH_RESERVED_DEVICE",
        r"C:\Projects\.": "CREATION_TARGET_PATH_SEGMENT_INVALID",
        "C:\\Projects\\道路工程\\": "CREATION_TARGET_PATH_SEGMENT_INVALID",
        r"C:\\Projects": "CREATION_TARGET_PATH_SEGMENT_INVALID",
    }
    for path, code in cases.items():
        codes = [item.code for item in validate_creation_target_path(path)]
        assert code in codes, (path, codes)
        assert all(item.sheet == "" and item.row is None for item in validate_creation_target_path(path))

    too_long = "C:\\Projects\\" + "道" * MAX_TARGET_PATH_LENGTH
    assert "CREATION_TARGET_PATH_TOO_LONG" in [
        item.code for item in validate_creation_target_path(too_long)
    ]
