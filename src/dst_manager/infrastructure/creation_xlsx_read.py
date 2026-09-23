"""创建模板工作簿的只读回读（PLAN-DM-036 Task 2 / SPEC-DM-018 §5）。

``read_creation_workbook`` 把工作簿字节读成只读快照（逐工作表文本网格、公式
单元格、无法解释的单元格、包内宏/外部链接部件），不做任何协议判定：工作表
集合、表头、隐藏元数据与标准身份是否合法由
:mod:`dst_manager.application.creation_import` 按固定标准判定。超出受控规模的
网格不展开，只报行列数，由调用方拒绝。

读取过程不执行公式、宏或外部链接；包内危险部件只做标记，由解析端拒绝。
"""

from __future__ import annotations

import zipfile
from io import BytesIO

from openpyxl import load_workbook

from dst_manager.infrastructure.creation_xlsx_protocol import (
    FORBIDDEN_PART_MARKERS,
    MAX_SHEET_COLUMNS,
    MAX_SHEET_ROWS,
    CreationSheetGrid,
    CreationWorkbookView,
    CreationXlsxReadError,
    column_letter,
)

__all__ = ["read_creation_workbook"]


def read_creation_workbook(data: bytes) -> CreationWorkbookView:
    """把工作簿字节读成只读快照；不执行公式、宏或外部链接。

    只做读取，不做协议判定；超出受控规模的网格不展开（``rows`` 为空）。
    无法读成 XLSX 时抛出 :class:`CreationXlsxReadError`。
    """
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            names = [name.lower() for name in archive.namelist()]
    except (OSError, zipfile.BadZipFile) as exc:
        raise CreationXlsxReadError(f"不是可读的 XLSX 包：{exc}") from exc
    try:
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:  # openpyxl 对损坏包抛出的异常类型不稳定
        raise CreationXlsxReadError(f"工作簿无法打开：{exc}") from exc
    try:
        formulas: list[tuple[str, int, str]] = []
        invalid: list[tuple[str, int, str]] = []
        sheets = tuple(_grid_of(sheet, formulas, invalid) for sheet in workbook.worksheets)
        return CreationWorkbookView(
            sheets=sheets,
            formulas=tuple(formulas),
            invalid_cells=tuple(invalid),
            external_links=bool(workbook._external_links),
            forbidden_parts=tuple(
                marker for marker in FORBIDDEN_PART_MARKERS if any(marker in name for name in names)
            ),
        )
    finally:
        workbook.close()


def _grid_of(
    worksheet,
    formulas: list[tuple[str, int, str]],
    invalid: list[tuple[str, int, str]],
) -> CreationSheetGrid:
    """读取一个工作表：超限只报行列数，否则展开文本网格并记录异常单元格。"""
    row_count = max(worksheet.max_row, 1)
    column_count = max(worksheet.max_column, 1)
    grid = CreationSheetGrid(
        name=worksheet.title,
        state=worksheet.sheet_state,
        row_count=row_count,
        column_count=column_count,
    )
    if row_count > MAX_SHEET_ROWS or column_count > MAX_SHEET_COLUMNS:
        return grid
    rows: list[tuple[str, ...]] = []
    for row_index, cells in enumerate(
        worksheet.iter_rows(min_row=1, max_row=row_count, min_col=1, max_col=column_count),
        start=1,
    ):
        values: list[str] = []
        for column_index, cell in enumerate(cells, start=1):
            letter = column_letter(column_index)
            if cell.data_type == "f":
                formulas.append((worksheet.title, row_index, letter))
            text = _cell_text(cell.value)
            if text is None:
                invalid.append((worksheet.title, row_index, letter))
                text = ""
            values.append(text)
        rows.append(tuple(values))
    return CreationSheetGrid(
        name=grid.name,
        state=grid.state,
        row_count=grid.row_count,
        column_count=grid.column_count,
        rows=tuple(rows),
    )


def _cell_text(value: object) -> str | None:
    """单元格值 → 文本：空值给空串，无法安全解释的类型给 ``None``。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, "g")
    return None
