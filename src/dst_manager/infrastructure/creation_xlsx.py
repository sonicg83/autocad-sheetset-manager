"""双工作表创建模板的生成与回读（PLAN-DM-036 Task 2 / SPEC-DM-018 §5）。

``build_creation_template`` 生成 legacy 风格的两张**可见**表：``SheetSet``
（A 列属性名、B 列输入值，另含固定「项目保存路径」行）与 ``Sheet``（第一行
固定表头 + 每行一个图纸组）。隐藏技术表 ``_CreationMeta`` 保存模板版本、精确
标准身份、各可见列的稳定 ``property_id``/协议键与资产 ``asset_id``；隐藏表
``_CreationLists`` 只放 Data Validation 候选值，因此候选值可以含逗号而不必
转义。可见表不出现派生属性列、内部 ID 或任何公式。

``read_creation_workbook`` 把工作簿字节读成只读快照（逐工作表文本网格、公式
单元格、无法解释的单元格、包内宏/外部链接部件），不做任何协议判定：工作表
集合、表头、隐藏元数据与标准身份是否合法由
:mod:`dst_manager.application.creation_import` 按固定标准判定。超出受控规模的
网格不展开，只报行列数，由调用方拒绝。

Data Validation 只是输入辅助，不是信任边界：解析端独立校验每个值。
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from dst_manager.domain.creation import (
    SHEET_SCOPE,
    SHEETSET_SCOPE,
    CreationAssetOption,
    ordinary_properties,
)
from dst_manager.domain.standard_models import (
    DrawingStandard,
    normalize_property_name,
)

__all__ = [
    "CREATION_TEMPLATE_VERSION",
    "FIXED_SHEET_COLUMNS",
    "HIDDEN_SHEET_NAMES",
    "LIST_SHEET",
    "MAX_GROUP_COUNT",
    "MAX_SHEET_COLUMNS",
    "MAX_SHEET_ROWS",
    "META_FIRST_RECORD_ROW",
    "META_HEADER",
    "META_SHEET",
    "SHEETSET_HEADERS",
    "SHEETSET_SHEET",
    "SHEET_SHEET",
    "TARGET_PATH_LABEL",
    "VALIDATION_ROW_LIMIT",
    "VISIBLE_SHEET_NAMES",
    "CreationAssetEntry",
    "CreationColumnSpec",
    "CreationSheetGrid",
    "CreationSheetSetRow",
    "CreationTemplatePlan",
    "CreationWorkbookView",
    "CreationXlsxReadError",
    "build_creation_template",
    "column_letter",
    "creation_metadata_records",
    "creation_template_plan",
    "read_creation_workbook",
]

#: 模板格式版本：隐藏技术表记录它，结构变化时必须递增。
CREATION_TEMPLATE_VERSION = 1

#: 两张可见工作表；顺序即生成顺序（``SheetSet`` 在前）。
SHEETSET_SHEET = "SheetSet"
SHEET_SHEET = "Sheet"
VISIBLE_SHEET_NAMES: tuple[str, str] = (SHEETSET_SHEET, SHEET_SHEET)
#: 隐藏技术表：元数据与 Data Validation 候选值。
META_SHEET = "_CreationMeta"
LIST_SHEET = "_CreationLists"
HIDDEN_SHEET_NAMES: tuple[str, str] = (META_SHEET, LIST_SHEET)

#: ``SheetSet`` 表：A 列属性名、B 列输入值；第 2 行固定为完整最终路径。
SHEETSET_HEADERS: tuple[str, str] = ("属性名", "输入值")
TARGET_PATH_LABEL = "项目保存路径"
#: ``Sheet`` 表的固定列（顺序即列序）：协议键 → 可见表头。
FIXED_SHEET_COLUMNS: tuple[tuple[str, str], ...] = (
    ("title", "图名"),
    ("count", "张数"),
    ("base_template", "基础模板"),
    ("layout_template", "布局模板"),
    ("paper_layout", "图幅"),
)
#: 动态列与固定表头重名时追加的可读限定语，保证可见表头唯一。
PROPERTY_HEADER_QUALIFIER = "（标准属性）"

#: 隐藏技术表的记录列与首条记录行。
META_HEADER: tuple[str, str, str, str] = ("record", "key", "value", "label")
META_FIRST_RECORD_ROW = 2
#: 隐藏技术表的记录类型。
META_RECORDS: frozenset[str] = frozenset(
    {
        "version",
        "standard",
        "fixed_column",
        "property_column",
        "sheetset_target_path",
        "sheetset_property",
        "asset",
    }
)

#: 受控规模：单个工作表的行号/列号上限，超出即整批拒绝。
MAX_SHEET_ROWS = 2000
MAX_SHEET_COLUMNS = 128
#: 单个图纸组的张数上限，避免异常工作簿触发放大。
MAX_GROUP_COUNT = 9999
#: Data Validation 覆盖的行上限（覆盖用户可能填写的范围）。
VALIDATION_ROW_LIMIT = 500

#: 包内禁止出现的部件标记：宏与外部链接。
_FORBIDDEN_PART_MARKERS = ("vbaproject", "externallink")

#: 样式：表头加粗居中，数据单元格顶端对齐 + 超长换行。
_HEADER_FONT = Font(bold=True)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_DATA_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
#: 列宽下限/上限（Excel 字符单位）与内容间距。
_COLUMN_WIDTH_PADDING = 2.0
_MIN_COLUMN_WIDTH = 10.0
_MAX_COLUMN_WIDTH = 40.0
#: CJK 区起点：全角字符按 2 个宽度单位估算（与图纸目录工作簿同口径）。
_CJK_WIDTH_THRESHOLD = 0x2E80


class CreationXlsxReadError(Exception):
    """工作簿字节无法读成 XLSX（非压缩包、包损坏或 openpyxl 打开失败）。"""


@dataclass(frozen=True, slots=True)
class CreationColumnSpec:
    """``Sheet`` 表的一列：列号、唯一可见表头与稳定映射键。

    ``key`` 对固定列是协议键（``title``/``count``/…），对动态列是标准属性
    ``property_id``；解析只按 ``key`` 定位，不按表头文案猜测身份。
    """

    index: int
    header: str
    key: str
    is_fixed: bool

    @property
    def column(self) -> str:
        return column_letter(self.index)


@dataclass(frozen=True, slots=True)
class CreationSheetSetRow:
    """``SheetSet`` 表的一行：行号、可见属性名与目标 ``property_id``。

    ``property_id`` 为 None 表示固定的「项目保存路径」行（完整最终路径不是
    标准属性）。
    """

    row: int
    label: str
    property_id: str | None = None

    @property
    def is_target_path(self) -> bool:
        return self.property_id is None


@dataclass(frozen=True, slots=True)
class CreationAssetEntry:
    """隐藏技术表里的一个资产标签映射：同类内唯一的用户标签 → 稳定 ``asset_id``。"""

    kind: str
    label: str
    asset_id: str


@dataclass(frozen=True, slots=True)
class CreationTemplatePlan:
    """模板的规范形状：列计划、``SheetSet`` 行计划与资产标签映射。"""

    columns: tuple[CreationColumnSpec, ...]
    sheetset_rows: tuple[CreationSheetSetRow, ...]
    assets: tuple[CreationAssetEntry, ...]
    #: 生成计划时的当前资产候选；「图幅」候选名从中去重提取。
    asset_options: tuple[CreationAssetOption, ...] = ()

    def column(self, key: str) -> CreationColumnSpec | None:
        for spec in self.columns:
            if spec.key == key:
                return spec
        return None

    def sheetset_row(self, property_id: str) -> CreationSheetSetRow | None:
        for row in self.sheetset_rows:
            if row.property_id == property_id:
                return row
        return None

    def asset_labels(self, kind: str) -> tuple[str, ...]:
        return tuple(entry.label for entry in self.assets if entry.kind == kind)

    def layout_names(self) -> tuple[str, ...]:
        """全部布局模板候选的布局名（去重后保持出现顺序），供「图幅」候选用。"""
        names: list[str] = []
        for option in self.asset_options:
            if option.kind != "layout-template":
                continue
            for name in option.layouts:
                if name not in names:
                    names.append(name)
        return tuple(names)


@dataclass(frozen=True, slots=True)
class CreationSheetGrid:
    """一个工作表的只读快照：行 × 列文本网格（空单元格为空串）。

    ``rows`` 为空表示该工作表超出受控规模、未展开；此时调用方按 ``row_count``
    与 ``column_count`` 拒绝。
    """

    name: str
    state: str
    row_count: int
    column_count: int
    rows: tuple[tuple[str, ...], ...] = ()


@dataclass(frozen=True, slots=True)
class CreationWorkbookView:
    """工作簿只读视图：工作表网格、公式单元格、不可解释单元格与包内危险部件。"""

    sheets: tuple[CreationSheetGrid, ...]
    formulas: tuple[tuple[str, int, str], ...] = ()
    invalid_cells: tuple[tuple[str, int, str], ...] = ()
    external_links: bool = False
    forbidden_parts: tuple[str, ...] = ()

    def sheet(self, name: str) -> CreationSheetGrid | None:
        for grid in self.sheets:
            if grid.name == name:
                return grid
        return None


def column_letter(index: int) -> str:
    """1 起始列号 → 列字母；诊断统一使用列字母。"""
    return get_column_letter(index)


def creation_template_plan(
    standard: DrawingStandard, asset_options: tuple[CreationAssetOption, ...]
) -> CreationTemplatePlan:
    """按标准与当前资产候选算出模板的规范形状。

    固定列在前、动态列按标准文档顺序在后；动态列表头与已用表头冲突时追加
    可读限定语，保证可见表头唯一。
    """
    columns: list[CreationColumnSpec] = []
    used: set[str] = set()
    for key, header in FIXED_SHEET_COLUMNS:
        used.add(normalize_property_name(header))
        columns.append(
            CreationColumnSpec(index=len(columns) + 1, header=header, key=key, is_fixed=True)
        )
    for prop in ordinary_properties(standard, SHEET_SCOPE):
        header = _unique_header(prop.name, used)
        used.add(normalize_property_name(header))
        columns.append(
            CreationColumnSpec(
                index=len(columns) + 1, header=header, key=prop.property_id, is_fixed=False
            )
        )

    sheetset_rows: list[CreationSheetSetRow] = [CreationSheetSetRow(row=2, label=TARGET_PATH_LABEL)]
    for prop in ordinary_properties(standard, SHEETSET_SCOPE):
        sheetset_rows.append(
            CreationSheetSetRow(
                row=len(sheetset_rows) + 2, label=prop.name, property_id=prop.property_id
            )
        )

    return CreationTemplatePlan(
        columns=tuple(columns),
        sheetset_rows=tuple(sheetset_rows),
        assets=tuple(
            CreationAssetEntry(kind=option.kind, label=option.label, asset_id=option.asset_id)
            for option in asset_options
        ),
        asset_options=tuple(asset_options),
    )


def _unique_header(name: str, used: set[str]) -> str:
    """动态列表头去重：与固定表头或其他属性名冲突时追加可读限定语。"""
    candidate = name
    index = 0
    while normalize_property_name(candidate) in used:
        index += 1
        qualifier = PROPERTY_HEADER_QUALIFIER if index == 1 else f"（标准属性 {index}）"
        candidate = f"{name}{qualifier}"
    return candidate


def creation_metadata_records(
    standard: DrawingStandard, plan: CreationTemplatePlan
) -> tuple[tuple[str, str, str, str], ...]:
    """隐藏技术表的规范记录序列（模板版本、标准身份、列/行映射与资产映射）。"""
    records: list[tuple[str, str, str, str]] = [
        ("version", "template_version", str(CREATION_TEMPLATE_VERSION), ""),
        ("standard", "standard_id", standard.standard_id, ""),
        ("standard", "standard_version", standard.version, ""),
    ]
    for spec in plan.columns:
        record = "fixed_column" if spec.is_fixed else "property_column"
        records.append((record, spec.column, spec.key, spec.header))
    for row in plan.sheetset_rows:
        if row.is_target_path:
            records.append(("sheetset_target_path", str(row.row), "", row.label))
        else:
            assert row.property_id is not None
            records.append(("sheetset_property", str(row.row), row.property_id, row.label))
    for entry in plan.assets:
        records.append(("asset", entry.kind, entry.asset_id, entry.label))
    return tuple(records)


# ---- 模板生成 ------------------------------------------------------------


def build_creation_template(
    standard: DrawingStandard, asset_options: tuple[CreationAssetOption, ...]
) -> bytes:
    """生成创建模板字节：两张可见表 + 隐藏技术表与候选值表 + Data Validation。"""
    plan = creation_template_plan(standard, asset_options)
    workbook = Workbook()
    sheetset_sheet = workbook.active
    sheetset_sheet.title = SHEETSET_SHEET
    sheet_sheet = workbook.create_sheet(SHEET_SHEET)
    meta_sheet = workbook.create_sheet(META_SHEET)
    meta_sheet.sheet_state = "hidden"
    list_sheet = workbook.create_sheet(LIST_SHEET)
    list_sheet.sheet_state = "hidden"

    _write_sheetset_sheet(sheetset_sheet, standard, plan)
    _write_sheet_sheet(sheet_sheet, plan)
    _write_metadata_sheet(meta_sheet, creation_metadata_records(standard, plan))
    candidates = _write_candidate_lists(list_sheet, standard, plan)
    _write_data_validations(sheetset_sheet, sheet_sheet, standard, plan, candidates)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _text_cell(worksheet, row: int, column: int, value: str, alignment: Alignment):
    """写入文本单元格：非空值固定为字符串类型（``=`` 开头压回文本，绝不成公式）。"""
    cell = worksheet.cell(row=row, column=column)
    if value != "":
        cell.value = value
        cell.data_type = "s"
    cell.alignment = alignment
    return cell


def _write_header_cell(worksheet, row: int, column: int, value: str) -> None:
    cell = _text_cell(worksheet, row, column, value, _HEADER_ALIGNMENT)
    cell.font = _HEADER_FONT


def _set_widths(worksheet, *columns: tuple[str, ...]) -> None:
    """逐列按内容最长显示宽度设列宽（夹在上下限之间）。"""
    for index, values in enumerate(columns, start=1):
        content = max((_display_width(value) for value in values), default=0)
        width = min(content + _COLUMN_WIDTH_PADDING, _MAX_COLUMN_WIDTH)
        worksheet.column_dimensions[column_letter(index)].width = max(width, _MIN_COLUMN_WIDTH)


def _display_width(text: str) -> int:
    """可读显示宽度估算：CJK 全角字符计 2，其余计 1。"""
    return sum(2 if ord(character) >= _CJK_WIDTH_THRESHOLD else 1 for character in text)


def _write_sheetset_sheet(
    worksheet, standard: DrawingStandard, plan: CreationTemplatePlan
) -> None:
    _write_header_cell(worksheet, 1, 1, SHEETSET_HEADERS[0])
    _write_header_cell(worksheet, 1, 2, SHEETSET_HEADERS[1])
    labels: list[str] = []
    values: list[str] = []
    for row_spec in plan.sheetset_rows:
        _text_cell(worksheet, row_spec.row, 1, row_spec.label, _DATA_ALIGNMENT)
        labels.append(row_spec.label)
        value = ""
        if row_spec.property_id is not None:
            # 模板预填标准默认值：导出再导入与界面初建得到同一份输入。
            value = standard.property_by_id(row_spec.property_id).default_value
        _text_cell(worksheet, row_spec.row, 2, value, _DATA_ALIGNMENT)
        values.append(value)
    _set_widths(worksheet, (SHEETSET_HEADERS[0], *labels), (SHEETSET_HEADERS[1], *values))
    worksheet.freeze_panes = "A2"


def _write_sheet_sheet(worksheet, plan: CreationTemplatePlan) -> None:
    for spec in plan.columns:
        _write_header_cell(worksheet, 1, spec.index, spec.header)
    _set_widths(worksheet, *((spec.header,) for spec in plan.columns))
    worksheet.freeze_panes = "A2"


def _write_metadata_sheet(
    worksheet, records: tuple[tuple[str, str, str, str], ...]
) -> None:
    for index, value in enumerate(META_HEADER, start=1):
        _write_header_cell(worksheet, 1, index, value)
    for offset, record in enumerate(records):
        for index, value in enumerate(record, start=1):
            _text_cell(worksheet, META_FIRST_RECORD_ROW + offset, index, value, _DATA_ALIGNMENT)
    _set_widths(
        worksheet,
        *((META_HEADER[index], *(record[index] for record in records)) for index in range(4)),
    )


def _candidate_columns(
    standard: DrawingStandard, plan: CreationTemplatePlan
) -> tuple[tuple[str, tuple[str, ...]], ...]:
    """Data Validation 候选列：稳定键 → 候选值（空候选列不写）。"""
    entries: list[tuple[str, tuple[str, ...]]] = []
    for prop in standard.properties:
        if prop.kind == "enum":
            entries.append((prop.property_id, tuple(item.value for item in prop.enum_items)))
    for key, kind in (("base_template", "base-template"), ("layout_template", "layout-template")):
        entries.append((key, plan.asset_labels(kind)))
    entries.append(("paper_layout", plan.layout_names()))
    return tuple(entry for entry in entries if entry[1])


def _write_candidate_lists(
    worksheet, standard: DrawingStandard, plan: CreationTemplatePlan
) -> dict[str, tuple[int, int]]:
    """把候选值写进隐藏表，返回 ``稳定键 → (列号, 候选个数)``。"""
    ranges: dict[str, tuple[int, int]] = {}
    for index, (key, values) in enumerate(_candidate_columns(standard, plan), start=1):
        _write_header_cell(worksheet, 1, index, key)
        for offset, value in enumerate(values):
            _text_cell(worksheet, 2 + offset, index, value, _DATA_ALIGNMENT)
        ranges[key] = (index, len(values))
    return ranges


def _write_data_validations(
    sheetset_sheet,
    sheet_sheet,
    standard: DrawingStandard,
    plan: CreationTemplatePlan,
    candidates: dict[str, tuple[int, int]],
) -> None:
    """枚举与三类模板/布局选项加 Data Validation；候选值引用隐藏表，避免转义问题。"""
    for prop in ordinary_properties(standard, SHEETSET_SCOPE):
        reference = _list_reference(candidates, prop.property_id)
        if reference is None:
            continue
        row = plan.sheetset_row(prop.property_id)
        assert row is not None
        _add_list_validation(sheetset_sheet, f"B{row.row}", reference)
    for spec in plan.columns:
        if spec.is_fixed:
            if spec.key not in ("base_template", "layout_template", "paper_layout"):
                continue
        elif standard.property_by_id(spec.key).kind != "enum":
            continue
        reference = _list_reference(candidates, spec.key)
        if reference is None:
            continue
        _add_list_validation(
            sheet_sheet, f"{spec.column}2:{spec.column}{VALIDATION_ROW_LIMIT}", reference
        )


def _list_reference(candidates: dict[str, tuple[int, int]], key: str) -> str | None:
    if key not in candidates:
        return None
    index, count = candidates[key]
    letter = column_letter(index)
    return f"{LIST_SHEET}!${letter}$2:${letter}${count + 1}"


def _add_list_validation(worksheet, cell_range: str, reference: str) -> None:
    validation = DataValidation(type="list", formula1=reference, allow_blank=True)
    worksheet.add_data_validation(validation)
    validation.add(cell_range)


# ---- 工作簿回读 ----------------------------------------------------------


def read_creation_workbook(data: bytes) -> CreationWorkbookView:
    """把工作簿字节读成只读快照；不执行公式、宏或外部链接。

    只做读取，不做协议判定；超出受控规模的网格不展开（``rows`` 为空）。
    无法读成 XLSX 时抛出 :class:`CreationXlsxReadError`。
    """
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            names = [name.lower() for name in archive.namelist()]
    except (OSError, zipfile.BadZipFile) as exc:
        raise CreationXlsxReadError(f"不是可读的 XLSX 包：{exc}") from exc
    try:
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:  # openpyxl 对损坏包抛出的异常类型不稳定
        raise CreationXlsxReadError(f"工作簿无法打开：{exc}") from exc
    try:
        formulas: list[tuple[str, int, str]] = []
        invalid: list[tuple[str, int, str]] = []
        sheets = tuple(_grid_of(sheet, formulas, invalid) for sheet in workbook.worksheets)
        return CreationWorkbookView(
            sheets=sheets,
            formulas=tuple(formulas),
            invalid_cells=tuple(invalid),
            external_links=bool(workbook._external_links),
            forbidden_parts=tuple(
                marker for marker in _FORBIDDEN_PART_MARKERS if any(marker in name for name in names)
            ),
        )
    finally:
        workbook.close()


def _grid_of(
    worksheet,
    formulas: list[tuple[str, int, str]],
    invalid: list[tuple[str, int, str]],
) -> CreationSheetGrid:
    """读取一个工作表：超限只报行列数，否则展开文本网格并记录异常单元格。"""
    row_count = max(worksheet.max_row, 1)
    column_count = max(worksheet.max_column, 1)
    grid = CreationSheetGrid(
        name=worksheet.title,
        state=worksheet.sheet_state,
        row_count=row_count,
        column_count=column_count,
    )
    if row_count > MAX_SHEET_ROWS or column_count > MAX_SHEET_COLUMNS:
        return grid
    rows: list[tuple[str, ...]] = []
    for row_index, cells in enumerate(
        worksheet.iter_rows(min_row=1, max_row=row_count, min_col=1, max_col=column_count),
        start=1,
    ):
        values: list[str] = []
        for column_index, cell in enumerate(cells, start=1):
            letter = column_letter(column_index)
            if cell.data_type == "f":
                formulas.append((worksheet.title, row_index, letter))
            text = _cell_text(cell.value)
            if text is None:
                invalid.append((worksheet.title, row_index, letter))
                text = ""
            values.append(text)
        rows.append(tuple(values))
    return CreationSheetGrid(
        name=grid.name,
        state=grid.state,
        row_count=grid.row_count,
        column_count=grid.column_count,
        rows=tuple(rows),
    )


def _cell_text(value: object) -> str | None:
    """单元格值 → 文本：空值给空串，无法安全解释的类型给 ``None``。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, "g")
    return None
