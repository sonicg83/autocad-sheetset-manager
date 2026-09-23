"""双工作表创建模板的生成（PLAN-DM-036 Task 2 / SPEC-DM-018 §5）。

``build_creation_template`` 生成 legacy 风格的两张**可见**表：``SheetSet``
（A 列属性名、B 列输入值，另含固定「项目保存路径」行）与 ``Sheet``（第一行
固定表头 + 每行一个图纸组），并把隐藏技术元数据与 Data Validation 候选值写入
隐藏表。可见表不出现派生属性列、内部 ID 或任何公式：单元格一律按文本写入
（``=`` 开头压回文本），因此模板本身不含公式。

模板形状取自 :mod:`dst_manager.infrastructure.creation_xlsx_protocol` 的规范计划；
Data Validation 只是输入辅助，不是信任边界，解析端独立校验每个值。
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from dst_manager.domain.creation import (
    SHEETSET_SCOPE,
    CreationAssetOption,
    ordinary_properties,
)
from dst_manager.domain.standard_models import DrawingStandard
from dst_manager.infrastructure.creation_xlsx_protocol import (
    LIST_SHEET,
    META_FIRST_RECORD_ROW,
    META_HEADER,
    META_SHEET,
    SHEET_SHEET,
    SHEETSET_HEADERS,
    SHEETSET_SHEET,
    VALIDATION_ROW_LIMIT,
    CreationTemplatePlan,
    column_letter,
    creation_metadata_records,
    creation_template_plan,
)

__all__ = ["build_creation_template"]

#: 样式：表头加粗居中，数据单元格顶端对齐 + 超长换行。
_HEADER_FONT = Font(bold=True)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_DATA_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
#: 列宽下限/上限（Excel 字符单位）与内容间距。
_COLUMN_WIDTH_PADDING = 2.0
_MIN_COLUMN_WIDTH = 10.0
_MAX_COLUMN_WIDTH = 40.0
#: CJK 区起点：全角字符按 2 个宽度单位估算（与图纸目录工作簿同口径）。
_CJK_WIDTH_THRESHOLD = 0x2E80


def build_creation_template(
    standard: DrawingStandard, asset_options: tuple[CreationAssetOption, ...]
) -> bytes:
    """生成创建模板字节：两张可见表 + 隐藏技术表与候选值表 + Data Validation。"""
    plan = creation_template_plan(standard, asset_options)
    workbook = Workbook()
    sheetset_sheet = workbook.active
    sheetset_sheet.title = SHEETSET_SHEET
    sheet_sheet = workbook.create_sheet(SHEET_SHEET)
    meta_sheet = workbook.create_sheet(META_SHEET)
    meta_sheet.sheet_state = "hidden"
    list_sheet = workbook.create_sheet(LIST_SHEET)
    list_sheet.sheet_state = "hidden"

    _write_sheetset_sheet(sheetset_sheet, standard, plan)
    _write_sheet_sheet(sheet_sheet, plan)
    _write_metadata_sheet(meta_sheet, creation_metadata_records(standard, plan))
    candidates = _write_candidate_lists(list_sheet, standard, plan)
    _write_data_validations(sheetset_sheet, sheet_sheet, standard, plan, candidates)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _text_cell(worksheet, row: int, column: int, value: str, alignment: Alignment):
    """写入文本单元格：非空值固定为字符串类型（``=`` 开头压回文本，绝不成公式）。"""
    cell = worksheet.cell(row=row, column=column)
    if value != "":
        cell.value = value
        cell.data_type = "s"
    cell.alignment = alignment
    return cell


def _write_header_cell(worksheet, row: int, column: int, value: str) -> None:
    cell = _text_cell(worksheet, row, column, value, _HEADER_ALIGNMENT)
    cell.font = _HEADER_FONT


def _set_widths(worksheet, *columns: tuple[str, ...]) -> None:
    """逐列按内容最长显示宽度设列宽（夹在上下限之间）。"""
    for index, values in enumerate(columns, start=1):
        content = max((_display_width(value) for value in values), default=0)
        width = min(content + _COLUMN_WIDTH_PADDING, _MAX_COLUMN_WIDTH)
        worksheet.column_dimensions[column_letter(index)].width = max(width, _MIN_COLUMN_WIDTH)


def _display_width(text: str) -> int:
    """可读显示宽度估算：CJK 全角字符计 2，其余计 1。"""
    return sum(2 if ord(character) >= _CJK_WIDTH_THRESHOLD else 1 for character in text)


def _write_sheetset_sheet(
    worksheet, standard: DrawingStandard, plan: CreationTemplatePlan
) -> None:
    _write_header_cell(worksheet, 1, 1, SHEETSET_HEADERS[0])
    _write_header_cell(worksheet, 1, 2, SHEETSET_HEADERS[1])
    labels: list[str] = []
    values: list[str] = []
    for row_spec in plan.sheetset_rows:
        _text_cell(worksheet, row_spec.row, 1, row_spec.label, _DATA_ALIGNMENT)
        labels.append(row_spec.label)
        value = ""
        if row_spec.property_id is not None:
            # 模板预填标准默认值：导出再导入与界面初建得到同一份输入。
            value = standard.property_by_id(row_spec.property_id).default_value
        _text_cell(worksheet, row_spec.row, 2, value, _DATA_ALIGNMENT)
        values.append(value)
    _set_widths(worksheet, (SHEETSET_HEADERS[0], *labels), (SHEETSET_HEADERS[1], *values))
    worksheet.freeze_panes = "A2"


def _write_sheet_sheet(worksheet, plan: CreationTemplatePlan) -> None:
    for spec in plan.columns:
        _write_header_cell(worksheet, 1, spec.index, spec.header)
    _set_widths(worksheet, *((spec.header,) for spec in plan.columns))
    worksheet.freeze_panes = "A2"


def _write_metadata_sheet(
    worksheet, records: tuple[tuple[str, str, str, str], ...]
) -> None:
    for index, value in enumerate(META_HEADER, start=1):
        _write_header_cell(worksheet, 1, index, value)
    for offset, record in enumerate(records):
        for index, value in enumerate(record, start=1):
            _text_cell(worksheet, META_FIRST_RECORD_ROW + offset, index, value, _DATA_ALIGNMENT)
    _set_widths(
        worksheet,
        *((META_HEADER[index], *(record[index] for record in records)) for index in range(4)),
    )


def _candidate_columns(
    standard: DrawingStandard, plan: CreationTemplatePlan
) -> tuple[tuple[str, tuple[str, ...]], ...]:
    """Data Validation 候选列：稳定键 → 候选值（空候选列不写）。"""
    entries: list[tuple[str, tuple[str, ...]]] = []
    for prop in standard.properties:
        if prop.kind == "enum":
            entries.append((prop.property_id, tuple(item.value for item in prop.enum_items)))
    for key, kind in (("base_template", "base-template"), ("layout_template", "layout-template")):
        entries.append((key, plan.asset_labels(kind)))
    entries.append(("paper_layout", plan.layout_names()))
    return tuple(entry for entry in entries if entry[1])


def _write_candidate_lists(
    worksheet, standard: DrawingStandard, plan: CreationTemplatePlan
) -> dict[str, tuple[int, int]]:
    """把候选值写进隐藏表，返回 ``稳定键 → (列号, 候选个数)``。"""
    ranges: dict[str, tuple[int, int]] = {}
    for index, (key, values) in enumerate(_candidate_columns(standard, plan), start=1):
        _write_header_cell(worksheet, 1, index, key)
        for offset, value in enumerate(values):
            _text_cell(worksheet, 2 + offset, index, value, _DATA_ALIGNMENT)
        ranges[key] = (index, len(values))
    return ranges


def _write_data_validations(
    sheetset_sheet,
    sheet_sheet,
    standard: DrawingStandard,
    plan: CreationTemplatePlan,
    candidates: dict[str, tuple[int, int]],
) -> None:
    """枚举与三类模板/布局选项加 Data Validation；候选值引用隐藏表，避免转义问题。"""
    for prop in ordinary_properties(standard, SHEETSET_SCOPE):
        reference = _list_reference(candidates, prop.property_id)
        if reference is None:
            continue
        row = plan.sheetset_row(prop.property_id)
        assert row is not None
        _add_list_validation(sheetset_sheet, f"B{row.row}", reference)
    for spec in plan.columns:
        if spec.is_fixed:
            if spec.key not in ("base_template", "layout_template", "paper_layout"):
                continue
        elif standard.property_by_id(spec.key).kind != "enum":
            continue
        reference = _list_reference(candidates, spec.key)
        if reference is None:
            continue
        _add_list_validation(
            sheet_sheet, f"{spec.column}2:{spec.column}{VALIDATION_ROW_LIMIT}", reference
        )


def _list_reference(candidates: dict[str, tuple[int, int]], key: str) -> str | None:
    if key not in candidates:
        return None
    index, count = candidates[key]
    letter = column_letter(index)
    return f"{LIST_SHEET}!${letter}$2:${letter}${count + 1}"


def _add_list_validation(worksheet, cell_range: str, reference: str) -> None:
    validation = DataValidation(type="list", formula1=reference, allow_blank=True)
    worksheet.add_data_validation(validation)
    validation.add(cell_range)
