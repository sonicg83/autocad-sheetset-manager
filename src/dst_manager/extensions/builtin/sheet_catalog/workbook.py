"""图纸目录候选工作簿生成与回读验证（PLAN-DM-020 Task 7 / SPEC-DM-012 §9）。

``write_candidate`` 由扩展侧调用：把规范化列（表头 + 全字符串数据行）写成
单工作表候选 XLSX，落盘后立即以 ``data_only=False/read_only=False`` 重开并
经 ``validate_candidate`` 自检，返回 ``WorkbookSummary``。
``validate_candidate`` 由宿主侧调用：重开候选文件核对工作表形状（唯一可见
"图纸目录"）、表头逐字一致、数据行数与禁止特性（公式/宏/图表/外部链接/
活动X/隐藏列/隐藏表/定义名称/超链接/非文本单元格），任何一条不满足都以
``SHEET_CATALOG_XLSX_INVALID`` 结构化报错。

样式全部受控常量化：不加公式、宏、图表、外部链接、隐藏列/表，也不嵌入
任何工程绝对路径（单元格一律字符串、无定义名称与超链接）。本模块不做
本地化、不读取翻译资源。
"""

from __future__ import annotations

import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from dst_manager.extensions.builtin.sheet_catalog.errors import (
    SheetCatalogError,
    sheet_catalog_error,
)

__all__ = [
    "MAX_COLUMN_WIDTH",
    "MIN_COLUMN_WIDTH",
    "WORKSHEET_NAME",
    "WorkbookSummary",
    "validate_candidate",
    "write_candidate",
]

#: 唯一允许的工作表名（同时固定 ``WorkbookSummary.worksheet_name`` 形状）。
WORKSHEET_NAME: Literal["图纸目录"] = "图纸目录"

#: 样式常量：表头加粗居中，数据单元格顶端对齐 + 超长换行。
_HEADER_FONT = Font(bold=True)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_DATA_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

#: 列宽下限/上限（Excel 字符单位）与内容间距；上限截断后依赖换行。
MIN_COLUMN_WIDTH = 8.0
MAX_COLUMN_WIDTH = 40.0
_COLUMN_WIDTH_PADDING = 2.0
#: CJK 区起点：全角字符按 2 个宽度单位估算。
_CJK_WIDTH_THRESHOLD = 0x2E80

#: XLSX 包内禁止出现的部件标记（宏/外部链接/图表/活动X）。
_FORBIDDEN_PART_MARKERS = (
    "vbaproject",
    "externallink",
    "chart",
    "activex",
)


@dataclass(frozen=True, slots=True)
class WorkbookSummary:
    """候选工作簿回读摘要：形状由 ``Literal`` 与元组钉住。"""

    worksheet_name: Literal["图纸目录"]
    headers: tuple[str, ...]
    data_rows: int


def _xlsx_invalid(check: str) -> SheetCatalogError:
    return sheet_catalog_error("SHEET_CATALOG_XLSX_INVALID", {"check": check})


def _display_width(text: str) -> int:
    """可读显示宽度估算：CJK 全角字符计 2，其余计 1。"""
    return sum(2 if ord(char) >= _CJK_WIDTH_THRESHOLD else 1 for char in text)


def _column_widths(
    headers: tuple[str, ...], rows: tuple[tuple[str, ...], ...]
) -> list[float]:
    """按表头与内容逐列取最大显示宽度，夹在上下限之间。"""
    widths: list[float] = []
    for index, header in enumerate(headers):
        values = [header]
        values.extend(row[index] for row in rows if index < len(row))
        content = max(_display_width(value) for value in values)
        width = min(content + _COLUMN_WIDTH_PADDING, MAX_COLUMN_WIDTH)
        widths.append(max(width, MIN_COLUMN_WIDTH))
    return widths


def _text_cell(
    worksheet, row: int, column: int, value: str, alignment: Alignment
):
    """写入单元格：非空值为字符串类型（``=`` 开头压回文本，绝不成为公式）。

    空值不赋内容但**仍创建单元格并施加样式**：XLSX 只落盘有样式的空单元格，
    这样行尾全空的数据行读回后行仍存在（``max_row`` 覆盖全部数据行），
    行数语义与输入一致；空单元格读回 ``None``/``n``。
    """
    cell = worksheet.cell(row=row, column=column)
    if value != "":
        cell.value = value
        cell.data_type = "s"
    cell.alignment = alignment
    return cell


def write_candidate(
    path: Path, headers: Sequence[str], rows: Iterable[Sequence[str]]
) -> WorkbookSummary:
    """写候选工作簿：单可见"图纸目录"表、首行列名、A2 冻结与自动筛选。"""
    header_tuple = tuple(str(header) for header in headers)
    row_tuples = tuple(tuple(str(value) for value in row) for row in rows)
    column_count = len(header_tuple)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = WORKSHEET_NAME
    worksheet.sheet_state = "visible"

    for column, header in enumerate(header_tuple, start=1):
        _text_cell(worksheet, 1, column, header, _HEADER_ALIGNMENT)
        worksheet.cell(row=1, column=column).font = _HEADER_FONT
    for row_index, row in enumerate(row_tuples, start=2):
        for column in range(column_count):
            value = row[column] if column < len(row) else ""
            # 空值也写占位单元格：保证全空/行尾全空数据行读回后行仍存在。
            _text_cell(worksheet, row_index, column + 1, value, _DATA_ALIGNMENT)

    worksheet.freeze_panes = "A2"
    last_row = max(1, len(row_tuples) + 1)
    last_column = get_column_letter(max(column_count, 1))
    worksheet.auto_filter.ref = f"A1:{last_column}{last_row}"

    for column, width in enumerate(_column_widths(header_tuple, row_tuples), start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width

    workbook.save(path)
    return validate_candidate(path, header_tuple, len(row_tuples))


def _check_forbidden_package_parts(path: Path) -> None:
    """候选包内出现宏/外部链接/图表/活动X部件即拒绝。"""
    try:
        with zipfile.ZipFile(path) as archive:
            names = [name.lower() for name in archive.namelist()]
    except (OSError, zipfile.BadZipFile) as exc:
        raise _xlsx_invalid("package_read") from exc
    for marker in _FORBIDDEN_PART_MARKERS:
        if any(marker in name for name in names):
            raise _xlsx_invalid(f"forbidden_part:{marker}")


def validate_candidate(
    path: Path, expected_headers: Sequence[str], expected_rows: int
) -> WorkbookSummary:
    """宿主回读校验：全部核对项通过才返回 ``WorkbookSummary``。"""
    expected = tuple(str(header) for header in expected_headers)
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        raise _xlsx_invalid("workbook_open") from exc
    try:
        _check_forbidden_package_parts(path)
        if len(workbook.sheetnames) != 1:
            raise _xlsx_invalid("worksheet_count")
        worksheet = workbook[workbook.sheetnames[0]]
        if worksheet.title != WORKSHEET_NAME:
            raise _xlsx_invalid("worksheet_name")
        if worksheet.sheet_state != "visible":
            raise _xlsx_invalid("worksheet_visible")
        if workbook._external_links:
            raise _xlsx_invalid("external_links")
        if len(workbook.defined_names) > 0:
            raise _xlsx_invalid("defined_names")

        header_row = tuple(
            str(cell.value) for cell in worksheet[1]
        )
        if header_row != expected:
            raise _xlsx_invalid("headers")

        data_rows = max(worksheet.max_row - 1, 0)
        if data_rows != expected_rows:
            raise _xlsx_invalid("data_row_count")

        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    raise _xlsx_invalid("formulas")
                if cell.data_type != "s":
                    raise _xlsx_invalid("non_text_cells")
                if cell.hyperlink is not None:
                    raise _xlsx_invalid("hyperlinks")

        for dimension in worksheet.column_dimensions.values():
            if dimension.hidden:
                raise _xlsx_invalid("hidden_columns")
    finally:
        workbook.close()
    return WorkbookSummary(
        worksheet_name=WORKSHEET_NAME,
        headers=expected,
        data_rows=expected_rows,
    )
